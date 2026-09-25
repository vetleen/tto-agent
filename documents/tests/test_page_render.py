"""Tests for documents.services.page_render (pptx slide renders via Gotenberg).

The render service is never contacted: ``GotenbergClient.convert`` is patched with a
fake that hands back a blank PDF with as many pages as the subset it received.
"""
from __future__ import annotations

import io
import tempfile
import types
from unittest.mock import patch
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import SimpleTestCase, TestCase, override_settings

from chat.models import Asset
from documents.models import DataRoom, DataRoomDocument, DataRoomDocumentVersion
from documents.services import page_render as pr
from documents.tests._helpers import make_document

User = get_user_model()

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _png(width: int = 8, height: int = 8, color=(200, 30, 30)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, format="PNG")
    return buf.getvalue()


def _jpeg(width: int = 16, height: int = 9) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), (10, 120, 200)).save(buf, format="JPEG")
    return buf.getvalue()


def _deck(n_slides: int, *, picture: bytes | None = None, slidenum_on: int | None = None,
          hide: int | None = None, notes: bool = True) -> bytes:
    """A python-pptx deck: slide i is titled ``Slide i``."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(1, n_slides + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = f"Slide {i}"
        if notes:
            slide.notes_slide.notes_text_frame.text = f"Notes {i}"
        if picture is not None:
            slide.shapes.add_picture(io.BytesIO(picture), Inches(1), Inches(2), Inches(3), Inches(2))
        if slidenum_on == i:
            from lxml import etree

            box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(0.5))
            paragraph = box.text_frame.paragraphs[0]._p
            fld = etree.SubElement(paragraph, f"{{{_A_NS}}}fld")
            fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
            fld.set("type", "slidenum")
            rpr = etree.SubElement(fld, f"{{{_A_NS}}}rPr")
            rpr.set("lang", "en-US")
            etree.SubElement(fld, f"{{{_A_NS}}}t").text = "‹#›"
        if hide == i:
            slide._element.set("show", "0")
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _blank_pdf(pages: int, width: float = 960, height: float = 540) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=width, height=height)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _slide_titles(pptx_bytes: bytes) -> list[str]:
    from pptx import Presentation

    return [s.shapes.title.text for s in Presentation(io.BytesIO(pptx_bytes)).slides]


class OptimizePptxTests(SimpleTestCase):
    def test_downscales_large_raster_keeping_its_format(self):
        from PIL import Image

        data = _deck(1, picture=_png(3000, 2000), notes=False)
        out = pr.optimize_pptx(data)
        self.assertLess(len(out), len(data))
        self.assertEqual(_slide_titles(out), ["Slide 1"])
        with ZipFile(io.BytesIO(out)) as z:
            media = [n for n in z.namelist() if n.startswith("ppt/media/")]
            self.assertEqual(len(media), 1)
            self.assertTrue(media[0].endswith(".png"), media[0])
            with Image.open(io.BytesIO(z.read(media[0]))) as img:
                self.assertEqual(img.format, "PNG")
                self.assertLessEqual(max(img.size), 1568)

    def test_strips_embedded_fonts_and_their_references(self):
        from chat.slides.pptx_fonts import embed_fonts

        face = types.SimpleNamespace(data=b"\x00\x01" * 64, weight=400, style="normal")
        data = embed_fonts(_deck(2, notes=False), [{"typeface": "Custom Sans", "faces": [face]}])
        with ZipFile(io.BytesIO(data)) as z:
            self.assertTrue(any(n.startswith("ppt/fonts/") for n in z.namelist()))

        out = pr.optimize_pptx(data)
        with ZipFile(io.BytesIO(out)) as z:
            names = z.namelist()
            self.assertFalse(any(n.startswith("ppt/fonts/") for n in names))
            pres = z.read("ppt/presentation.xml").decode("utf-8")
            self.assertNotIn("embeddedFontLst", pres)
            self.assertNotIn("embedTrueTypeFonts", pres)
            self.assertNotIn("relationships/font", z.read("ppt/_rels/presentation.xml.rels").decode("utf-8"))
            self.assertNotIn("fntdata", z.read("[Content_Types].xml").decode("utf-8"))
        self.assertEqual(_slide_titles(out), ["Slide 1", "Slide 2"])

    def test_returns_input_unchanged_on_garbage(self):
        self.assertEqual(pr.optimize_pptx(b"not a zip"), b"not a zip")


class SplitPptxTests(SimpleTestCase):
    def test_keeps_requested_slides_in_order_and_drops_their_notes(self):
        data = _deck(5)
        out = pr.split_pptx(data, [2, 4])
        self.assertEqual(_slide_titles(out), ["Slide 2", "Slide 4"])
        with ZipFile(io.BytesIO(out)) as z:
            names = z.namelist()
        self.assertEqual(len([n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]), 2)
        self.assertEqual(len([n for n in names if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")]), 2)
        # The original is untouched.
        self.assertEqual(pr.count_slides(data), 5)

    def test_literalizes_slide_number_fields_and_unhides(self):
        data = _deck(4, slidenum_on=3, hide=3, notes=False)
        out = pr.split_pptx(data, [3])
        with ZipFile(io.BytesIO(out)) as z:
            # python-pptx keeps the surviving part's original name (slide3.xml).
            slide_parts = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
            self.assertEqual(len(slide_parts), 1)
            slide_xml = z.read(slide_parts[0]).decode("utf-8")
        self.assertNotIn('type="slidenum"', slide_xml)
        self.assertIn("<a:t>3</a:t>", slide_xml)
        self.assertNotIn('show="0"', slide_xml)
        # The run keeps the field's run properties.
        self.assertIn('lang="en-US"', slide_xml)

    def test_single_slide_subset(self):
        out = pr.split_pptx(_deck(3, notes=False), [3])
        self.assertEqual(_slide_titles(out), ["Slide 3"])
        self.assertEqual(pr.count_slides(out), 1)


def _response(status: int, content: bytes):
    return types.SimpleNamespace(status_code=status, content=content)


class GotenbergClientTests(SimpleTestCase):
    def setUp(self):
        self.client_ = pr.GotenbergClient("http://render.test/", user="u", password="p", timeout=7)

    def _convert(self, post):
        import requests

        with patch.object(requests.Session, "post", post):
            return self.client_.convert(b"PK-fake", upload_name="12-1.pptx", trace="12")

    def test_success_returns_pdf_and_uses_opaque_name_and_trace(self):
        calls = {}

        def post(session, url, files=None, headers=None, timeout=None):
            calls.update(url=url, files=files, headers=headers, timeout=timeout)
            return _response(200, b"%PDF-1.7 fake")

        self.assertEqual(self._convert(post), b"%PDF-1.7 fake")
        self.assertEqual(calls["url"], "http://render.test/forms/libreoffice/convert")
        self.assertEqual(calls["files"]["files"][0], "12-1.pptx")
        self.assertEqual(calls["headers"]["Gotenberg-Trace"], "12")
        self.assertEqual(calls["timeout"], (10, 7))
        self.assertEqual(self.client_._session.auth, ("u", "p"))

    def test_status_mapping(self):
        cases = [
            (503, b"queue is full", pr.RenderBusy),
            (429, b"", pr.RenderBusy),
            (503, b"<!DOCTYPE html><html>herokucdn</html>", pr.RenderTimeout),
            (503, b"the request has timed out", pr.RenderTimeout),
            (504, b"", pr.RenderTimeout),
            (500, b"LibreOffice failed to convert", pr.RenderRejected),
            (400, b"", pr.RenderRejected),
            (200, b"not a pdf", pr.RenderRejected),
            (401, b"", pr.RenderUnavailable),
            (502, b"", pr.RenderUnavailable),
        ]
        for status, body, exc_type in cases:
            with self.subTest(status=status, body=body[:12]):
                with self.assertRaises(exc_type):
                    self._convert(lambda *a, **k: _response(status, body))

    def test_connection_and_timeout_errors(self):
        import requests

        def boom(*a, **k):
            raise requests.ConnectionError("refused")

        def slow(*a, **k):
            raise requests.Timeout("read timed out")

        with self.assertRaises(pr.RenderUnavailable):
            self._convert(boom)
        with self.assertRaises(pr.RenderTimeout):
            self._convert(slow)

    def test_client_from_settings_is_none_when_disabled(self):
        with override_settings(DOCUMENT_RENDER_SERVICE_URL=""):
            self.assertIsNone(pr.client_from_settings())
            self.assertFalse(pr.is_enabled())
        with override_settings(DOCUMENT_RENDER_SERVICE_URL="http://render.test"):
            self.assertIsNotNone(pr.client_from_settings())
            self.assertTrue(pr.is_enabled())


class RasterizeTests(SimpleTestCase):
    def test_jpeg_per_page_at_vision_cap(self):
        from PIL import Image

        jpegs = pr.rasterize_pdf(_blank_pdf(2), 2)
        self.assertEqual(len(jpegs), 2)
        with Image.open(io.BytesIO(jpegs[0])) as img:
            self.assertEqual(img.format, "JPEG")
            self.assertEqual(max(img.size), 1568)

    def test_page_count_mismatch_is_rejected(self):
        with self.assertRaises(pr.RenderRejected):
            pr.rasterize_pdf(_blank_pdf(2), 3)
        with self.assertRaises(pr.RenderRejected):
            pr.rasterize_pdf(b"garbage", 1)


class _FakeConvert:
    """Stands in for GotenbergClient.convert: returns a blank PDF matching the subset.

    ``fail`` is a callable ``(call_no, n_slides, titles) -> Exception | None``.
    """

    def __init__(self, fail=None):
        self.calls: list[tuple[int, list[str], str]] = []
        self.fail = fail

    def __call__(self, pptx_bytes, *, upload_name, trace):
        # Patched onto the class as a plain (non-function) callable, so no ``self``.
        titles = _slide_titles(pptx_bytes)
        self.calls.append((len(titles), titles, upload_name))
        if self.fail is not None:
            exc = self.fail(len(self.calls), len(titles), titles)
            if exc is not None:
                raise exc
        return _blank_pdf(len(titles))


_MEDIA = tempfile.mkdtemp()


@override_settings(
    MEDIA_ROOT=_MEDIA,
    DOCUMENT_RENDER_SERVICE_URL="http://render.test",
    DOCUMENT_RENDER_BATCH_SIZE=8,
    DOCUMENT_RENDER_MAX_SLIDES=200,
)
class RenderVersionPagesTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="render@example.com", password="pw")
        self.room = DataRoom.objects.create(name="R", slug="r", created_by=self.user)

    def _version(self, deck: bytes, **kwargs) -> DataRoomDocumentVersion:
        doc = make_document(self.room, self.user, original_filename="deck.pptx")
        version = doc.current_version
        version.parser_type = "pptx"
        for key, value in kwargs.items():
            setattr(version, key, value)
        version.save()
        version.native_blob.save("deck.pptx", ContentFile(deck), save=True)
        return version

    def _renders(self, version):
        return list(
            Asset.objects.filter(version=version, role=Asset.ROLE_PAGE_RENDER)
            .order_by("page_number").values_list("page_number", flat=True)
        )

    def test_happy_path_renders_every_slide_in_batches(self):
        version = self._version(_deck(11, notes=False))
        fake = _FakeConvert()
        with patch.object(pr.GotenbergClient, "convert", fake):
            state = pr.render_version_pages(version.id)
        self.assertEqual(state, "ready")
        version.refresh_from_db()
        self.assertEqual(version.page_render_state, "ready")
        self.assertEqual(version.page_count, 11)
        self.assertEqual(version.page_render_attempts, 1)
        self.assertEqual(version.page_render_error, "")
        self.assertEqual(self._renders(version), list(range(1, 12)))
        self.assertEqual([c[0] for c in fake.calls], [8, 3])
        self.assertEqual(fake.calls[0][1][:2], ["Slide 1", "Slide 2"])
        self.assertEqual(fake.calls[1][1], ["Slide 9", "Slide 10", "Slide 11"])
        # Opaque upload names: version id + batch number only.
        self.assertEqual(fake.calls[0][2], f"{version.id}-1.pptx")
        asset = Asset.objects.get(version=version, role=Asset.ROLE_PAGE_RENDER, page_number=5)
        self.assertEqual(asset.content_type, "image/jpeg")
        self.assertTrue(asset.blob)
        self.assertEqual(max(asset.width, asset.height), 1568)
        self.assertEqual(asset.kind, Asset.KIND_IMAGE)
        self.assertEqual(asset.description, "Slide 5")

    def test_deck_over_the_cap_is_skipped_without_contacting_the_service(self):
        version = self._version(_deck(6, notes=False))
        fake = _FakeConvert()
        with override_settings(DOCUMENT_RENDER_MAX_SLIDES=5), patch.object(pr.GotenbergClient, "convert", fake):
            state = pr.render_version_pages(version.id)
        self.assertEqual(state, "skipped")
        version.refresh_from_db()
        self.assertEqual(version.page_count, 6)
        self.assertIn("6 slides", version.page_render_error)
        self.assertEqual(fake.calls, [])
        self.assertEqual(self._renders(version), [])

    def test_rejected_batch_falls_back_to_single_slides_and_ends_partial(self):
        version = self._version(_deck(3, notes=False))

        def fail(call_no, n_slides, titles):
            if n_slides > 1:
                return pr.RenderRejected("batch failed")
            if titles == ["Slide 2"]:
                return pr.RenderTimeout("too slow")
            return None

        fake = _FakeConvert(fail)
        with patch.object(pr.GotenbergClient, "convert", fake):
            state = pr.render_version_pages(version.id)
        self.assertEqual(state, "partial")
        version.refresh_from_db()
        self.assertEqual(version.page_render_error, "Slides 2 could not be rendered.")
        self.assertEqual(self._renders(version), [1, 3])
        self.assertEqual([c[0] for c in fake.calls], [3, 1, 1, 1])
        self.assertEqual(fake.calls[2][2], f"{version.id}-1-2.pptx")

    def test_every_slide_failing_ends_failed(self):
        version = self._version(_deck(2, notes=False))
        fake = _FakeConvert(lambda *_: pr.RenderRejected("nope"))
        with patch.object(pr.GotenbergClient, "convert", fake):
            state = pr.render_version_pages(version.id)
        self.assertEqual(state, "failed")
        version.refresh_from_db()
        self.assertEqual(version.page_render_error, "No slides could be rendered.")
        self.assertEqual(self._renders(version), [])

    def test_resume_only_renders_missing_slides(self):
        version = self._version(_deck(5, notes=False))
        pr.store_page_render(version, 1, _jpeg())
        pr.store_page_render(version, 2, _jpeg())
        fake = _FakeConvert()
        with patch.object(pr.GotenbergClient, "convert", fake):
            state = pr.render_version_pages(version.id)
        self.assertEqual(state, "ready")
        self.assertEqual([c[1] for c in fake.calls], [["Slide 3", "Slide 4", "Slide 5"]])
        self.assertEqual(self._renders(version), [1, 2, 3, 4, 5])

    def test_store_page_render_is_idempotent(self):
        version = self._version(_deck(1, notes=False))
        first = pr.store_page_render(version, 1, _jpeg())
        second = pr.store_page_render(version, 1, _jpeg(20, 20))
        self.assertEqual(first.id, second.id)
        self.assertEqual(Asset.objects.filter(version=version).count(), 1)
        self.assertEqual(pr.existing_page_numbers(version), {1})

    def test_disabled_service_marks_failed_without_http(self):
        version = self._version(_deck(1, notes=False))
        fake = _FakeConvert()
        with override_settings(DOCUMENT_RENDER_SERVICE_URL=""), patch.object(pr.GotenbergClient, "convert", fake):
            state = pr.render_version_pages(version.id)
        self.assertEqual(state, "failed")
        version.refresh_from_db()
        self.assertIn("not configured", version.page_render_error)
        self.assertEqual(fake.calls, [])

    def test_quarantined_or_non_pptx_versions_are_skipped(self):
        quarantined = self._version(_deck(1, notes=False), is_quarantined=True)
        docx = self._version(_deck(1, notes=False), parser_type="docx")
        with patch.object(pr.GotenbergClient, "convert", _FakeConvert()):
            self.assertEqual(pr.render_version_pages(quarantined.id), "skipped")
            self.assertEqual(pr.render_version_pages(docx.id), "skipped")
        self.assertFalse(pr.should_render(quarantined))

    def test_busy_service_retries_in_place_then_propagates_leaving_pending(self):
        version = self._version(_deck(2, notes=False))
        fake = _FakeConvert(lambda *_: pr.RenderBusy("full"))
        with patch.object(pr.GotenbergClient, "convert", fake), patch.object(pr.time, "sleep") as sleep:
            with self.assertRaises(pr.RenderBusy):
                pr.render_version_pages(version.id)
        self.assertEqual(len(fake.calls), 1 + len(pr._BUSY_RETRY_DELAYS_S))
        self.assertEqual([c.args[0] for c in sleep.call_args_list], list(pr._BUSY_RETRY_DELAYS_S))
        version.refresh_from_db()
        self.assertEqual(version.page_render_state, "pending")
        self.assertEqual(version.page_render_attempts, 1)

    def test_unavailable_service_propagates_and_keeps_progress(self):
        version = self._version(_deck(10, notes=False))

        def fail(call_no, n_slides, titles):
            return pr.RenderUnavailable("down") if call_no == 2 else None

        fake = _FakeConvert(fail)
        with patch.object(pr.GotenbergClient, "convert", fake):
            with self.assertRaises(pr.RenderUnavailable):
                pr.render_version_pages(version.id)
        version.refresh_from_db()
        self.assertEqual(version.page_render_state, "pending")
        self.assertEqual(self._renders(version), list(range(1, 9)))
        # A retried delivery must not spend another attempt.
        with patch.object(pr.GotenbergClient, "convert", _FakeConvert()):
            self.assertEqual(pr.render_version_pages(version.id, count_attempt=False), "ready")
        version.refresh_from_db()
        self.assertEqual(version.page_render_attempts, 1)
        self.assertEqual(self._renders(version), list(range(1, 11)))

    def test_missing_version_is_a_noop(self):
        self.assertEqual(pr.render_version_pages(999_999), "none")


@override_settings(MEDIA_ROOT=_MEDIA, DOCUMENT_RENDER_SERVICE_URL="http://render.test")
class EnqueuePageRenderTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="enqueue@example.com", password="pw")
        self.room = DataRoom.objects.create(name="E", slug="e", created_by=self.user)
        doc = make_document(self.room, self.user, original_filename="deck.pptx")
        self.version = doc.current_version
        self.version.parser_type = "pptx"
        self.version.page_render_error = "old"
        self.version.save()

    def test_marks_pending_and_dispatches(self):
        with patch("documents.tasks.render_document_pages.delay") as delay:
            self.assertTrue(pr.enqueue_page_render(self.version.id))
        delay.assert_called_once_with(self.version.id)
        self.version.refresh_from_db()
        self.assertEqual(self.version.page_render_state, "pending")
        self.assertEqual(self.version.page_render_error, "")

    def test_dispatch_failure_leaves_pending_for_the_sweeper(self):
        with patch("documents.tasks.render_document_pages.delay", side_effect=RuntimeError("broker down")):
            self.assertFalse(pr.enqueue_page_render(self.version.id))
        self.version.refresh_from_db()
        self.assertEqual(self.version.page_render_state, "pending")

    def test_should_render_requires_ready_pptx_with_service_configured(self):
        self.assertTrue(pr.should_render(self.version))
        self.version.status = DataRoomDocument.Status.SCANNING
        self.assertFalse(pr.should_render(self.version))
        self.version.status = DataRoomDocument.Status.READY
        with override_settings(DOCUMENT_RENDER_SERVICE_URL=""):
            self.assertFalse(pr.should_render(self.version))
