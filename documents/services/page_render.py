"""Per-slide image renders for data-room ``.pptx`` versions.

Users upload a deck and expect the assistant to *see* it; text extraction loses the
layout. After a pptx version is released (READY, not quarantined) this module renders
every slide to a JPEG through the external **Gotenberg** app (LibreOffice behind HTTP —
see ``deploy/gotenberg/``) and stores the images as version-owned ``chat.Asset`` rows
with ``role=page_render``. ``document_view_native`` then serves selected slides to
the model. Page renders are never tokenized (no ``[[image:uuid]]``) and never shown
to users.

Pipeline per version (``render_version_pages``):

1. ``optimize_pptx`` — one zip rewrite: rasters downscaled to the vision cap (same
   format, so no part renames) and embedded fonts stripped properly (parts, rels,
   ``<p:embeddedFontLst>``, content types). Only ever sent to the render service,
   never stored.
2. ``split_pptx`` — python-pptx subsets of ``DOCUMENT_RENDER_BATCH_SIZE`` slides.
   LibreOffice's time and memory scale with the slides it *loads* (page ranges
   don't help), so batches bound the render dyno's RSS and keep every request far
   under Heroku's 30 s router limit. Slide-number fields are pinned to literals so
   they don't renumber inside a subset.
3. ``GotenbergClient.convert`` — POST the subset, get a PDF back.
4. ``rasterize_pdf`` — pypdfium2 → JPEG per page at the vision cap.
5. ``store_page_render`` — one Asset per slide, unique per (version, page_number),
   so a retried task resumes at the first missing slide.

A batch the service rejects is retried slide by slide (one bad slide only loses
itself). Busy/unavailable errors propagate so Celery retries with backoff; the
version stays ``pending`` and the stale sweeper bounds the attempts.

Only one conversion is in flight per worker process (``_render_slot``); the render
app's own queue limit (503 when full) is the cross-dyno backstop. Uploads use
opaque names (``<version_id>-<batch>.pptx``) — the render service's logs never see
a document name. Nothing here logs document names or text either.
"""
from __future__ import annotations

import copy
import hashlib
import io
import logging
import threading
import time
from contextlib import contextmanager
from zipfile import ZIP_DEFLATED, ZipFile

from django.conf import settings
from django.db import IntegrityError
from django.db.models import F
from django.utils import timezone

from documents.models import DataRoomDocument, DataRoomDocumentVersion

logger = logging.getLogger(__name__)

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PPTX_PARSER_TYPE = "pptx"

# OOXML namespaces / package parts (mirrors chat/slides/pptx_fonts.py, the inverse
# operation).
_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_FONT_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
_PRESENTATION = "ppt/presentation.xml"
_PRES_RELS = "ppt/_rels/presentation.xml.rels"
_CONTENT_TYPES = "[Content_Types].xml"
_FONT_PREFIX = "ppt/fonts/"
_MEDIA_PREFIX = "ppt/media/"
_RASTER_EXTS = {"png", "jpg", "jpeg", "gif", "webp"}

# In-place retries when the render service reports it is busy (queue full):
# sleep this long between attempts before giving the batch back to Celery.
_BUSY_RETRY_DELAYS_S = (5, 15, 30)
# While waiting for the render slot, heartbeat ``updated_at`` this often so the
# stale sweeper (STALE_PAGE_RENDER_MINUTES) leaves a merely-queued render alone.
_SLOT_WAIT_HEARTBEAT_S = 60.0

_render_semaphore = threading.BoundedSemaphore(1)


class RenderError(RuntimeError):
    """Base class for render-service failures."""


class RenderBusy(RenderError):
    """The service is saturated (queue full / rate limited) — try again later."""


class RenderTimeout(RenderError):
    """The conversion took too long (service timeout or the router's 30 s cut).

    Treated like a rejection: a smaller batch is the fix, not a later retry.
    """


class RenderRejected(RenderError):
    """The service could not convert this input (never retried as-is)."""


class RenderUnavailable(RenderError):
    """Configuration or connectivity problem (auth, DNS, connection refused)."""


# --------------------------------------------------------------------------- settings


def service_url() -> str:
    return (getattr(settings, "DOCUMENT_RENDER_SERVICE_URL", "") or "").strip().rstrip("/")


def is_enabled() -> bool:
    """Feature switch: an empty ``DOCUMENT_RENDER_SERVICE_URL`` disables rendering."""
    return bool(service_url())


def _int_setting(name: str, default: int, *, floor: int = 1) -> int:
    try:
        return max(floor, int(getattr(settings, name, default)))
    except (TypeError, ValueError):
        return default


# --------------------------------------------------------------------------- pptx prep


def optimize_pptx(data: bytes) -> bytes:
    """Shrink a deck for conversion: downscale rasters, drop embedded fonts.

    Best-effort — any failure logs a warning and returns the input unchanged.
    The output is a transient artefact for the render service only.
    """
    try:
        from lxml import etree

        with ZipFile(io.BytesIO(data)) as zin:
            entries = {name: zin.read(name) for name in zin.namelist()}
        _downscale_media(entries)
        _strip_embedded_fonts(entries, etree)
        out = io.BytesIO()
        with ZipFile(out, "w", ZIP_DEFLATED) as zout:
            for name, blob in entries.items():
                zout.writestr(name, blob)
        return out.getvalue()
    except Exception:  # noqa: BLE001 — optimisation is optional
        logger.warning("optimize_pptx: optimisation failed; converting the original bytes", exc_info=True)
        return data


def _downscale_media(entries: dict[str, bytes]) -> None:
    """Downscale raster media in place, keeping each part's format (no renames,
    so relationships and content types stay valid)."""
    from core.images import optimize_for_vision

    for name in list(entries):
        if not name.startswith(_MEDIA_PREFIX):
            continue
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        if ext not in _RASTER_EXTS:
            continue
        result = optimize_for_vision(entries[name], allow_transcode=False)
        if result is None:
            continue
        new_bytes, _media_type = result
        if len(new_bytes) < len(entries[name]):
            entries[name] = new_bytes


def _strip_embedded_fonts(entries: dict[str, bytes], etree) -> None:
    """Remove embedded font parts and every reference to them.

    Deleting only the ``.fntdata`` parts leaves dangling relationships that make
    LibreOffice fail to open the package; the presentation part, its rels and the
    content types must lose their entries too.
    """
    font_parts = [name for name in entries if name.startswith(_FONT_PREFIX)]
    if not font_parts:
        return
    for name in font_parts:
        del entries[name]

    if _PRES_RELS in entries:
        rels = etree.fromstring(entries[_PRES_RELS])
        for rel in list(rels):
            if rel.get("Type") == _FONT_REL_TYPE:
                rels.remove(rel)
        entries[_PRES_RELS] = _serialize(etree, rels)

    if _PRESENTATION in entries:
        pres = etree.fromstring(entries[_PRESENTATION])
        for element in pres.findall(f"{{{_P_NS}}}embeddedFontLst"):
            pres.remove(element)
        pres.attrib.pop("embedTrueTypeFonts", None)
        pres.attrib.pop("saveSubsetFonts", None)
        entries[_PRESENTATION] = _serialize(etree, pres)

    if _CONTENT_TYPES in entries:
        cts = etree.fromstring(entries[_CONTENT_TYPES])
        for element in list(cts):
            if element.tag == f"{{{_CT_NS}}}Default" and (element.get("Extension") or "").lower() == "fntdata":
                cts.remove(element)
            elif element.tag == f"{{{_CT_NS}}}Override" and (element.get("PartName") or "").startswith("/" + _FONT_PREFIX):
                cts.remove(element)
        entries[_CONTENT_TYPES] = _serialize(etree, cts)


def _serialize(etree, element) -> bytes:
    return etree.tostring(element, xml_declaration=True, encoding="UTF-8", standalone=True)


def count_slides(data: bytes) -> int:
    from pptx import Presentation

    return len(Presentation(io.BytesIO(data)).slides)


def literalize_slide_numbers(slide_element, number: int) -> None:
    """Replace ``<a:fld type="slidenum">`` runs with a plain run holding ``number``.

    Only changing the field's cached text is not enough: LibreOffice re-evaluates
    the field to the slide's index within the (sub)deck it is rendering.
    """
    from lxml import etree

    fld_tag = f"{{{_A_NS}}}fld"
    for fld in list(slide_element.iter(fld_tag)):
        if fld.get("type") != "slidenum":
            continue
        parent = fld.getparent()
        if parent is None:
            continue
        run = etree.Element(f"{{{_A_NS}}}r")
        rpr = fld.find(f"{{{_A_NS}}}rPr")
        if rpr is not None:
            run.append(copy.deepcopy(rpr))
        text = etree.SubElement(run, f"{{{_A_NS}}}t")
        text.text = str(number)
        parent.insert(parent.index(fld), run)
        parent.remove(fld)


def split_pptx(data: bytes, slide_numbers: list[int]) -> bytes:
    """Return a copy of the deck containing only the given 1-based slides, in order.

    Dropped slides (and the notes/media only they referenced) are not written,
    because python-pptx serialises only parts reachable from the package. Kept
    slides are un-hidden (LibreOffice skips hidden slides when exporting, which
    would shift the page ↔ slide mapping) and get literal slide numbers.
    """
    from pptx import Presentation

    keep = set(slide_numbers)
    prs = Presentation(io.BytesIO(data))
    sld_id_lst = prs.slides._sldIdLst
    slides = list(prs.slides)
    for idx, (sld_id, slide) in enumerate(zip(list(sld_id_lst.sldId_lst), slides), start=1):
        if idx not in keep:
            # Remove the XML reference FIRST: drop_rel only drops a relationship
            # whose reference count is already below 2.
            sld_id_lst.remove(sld_id)
            prs.part.drop_rel(sld_id.rId)
            continue
        slide._element.attrib.pop("show", None)
        literalize_slide_numbers(slide._element, idx)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- service


class GotenbergClient:
    """Thin client for Gotenberg's LibreOffice route (``/forms/libreoffice/convert``)."""

    def __init__(self, base_url: str, user: str = "", password: str = "", timeout: int = 60):
        import requests

        self._base = base_url.rstrip("/")
        self._timeout = max(1, int(timeout))
        self._session = requests.Session()
        if user or password:
            self._session.auth = (user, password)

    def convert(self, pptx_bytes: bytes, *, upload_name: str, trace: str) -> bytes:
        """Convert ``pptx_bytes`` to PDF bytes; raises a ``RenderError`` subclass."""
        import requests

        url = f"{self._base}/forms/libreoffice/convert"
        headers = {"Gotenberg-Trace": str(trace), "Gotenberg-Output-Filename": "out"}
        started = time.monotonic()
        try:
            response = self._session.post(
                url,
                files={"files": (upload_name, io.BytesIO(pptx_bytes), PPTX_MIME)},
                headers=headers,
                timeout=(10, self._timeout),
            )
        except requests.Timeout as exc:
            raise RenderTimeout(f"render service timed out after {self._timeout}s") from exc
        except requests.RequestException as exc:
            raise RenderUnavailable(f"render service unreachable ({exc.__class__.__name__})") from exc

        elapsed = time.monotonic() - started
        status = response.status_code
        body = response.content or b""
        if status == 200 and body[:5] == b"%PDF-":
            logger.info(
                "page_render: converted trace=%s bytes_in=%s bytes_out=%s in %.1fs",
                trace, len(pptx_bytes), len(body), elapsed,
            )
            return body

        logger.warning(
            "page_render: convert failed status=%s trace=%s bytes_in=%s in %.1fs",
            status, trace, len(pptx_bytes), elapsed,
        )
        raise _classify_failure(status, body, self._timeout)


def _classify_failure(status: int, body: bytes, timeout: int) -> RenderError:
    head = body[:512].lstrip().lower()
    if status == 200:
        return RenderRejected("render service returned a non-PDF response")
    if status == 504:
        return RenderTimeout("render service reported a gateway timeout")
    if status == 503:
        # Heroku's router answers an H12 (30 s) cut with its HTML error page;
        # Gotenberg's own 503s are plain text ("queue is full" / timeout).
        if head.startswith(b"<!doctype") or head.startswith(b"<html") or b"herokucdn" in head:
            return RenderTimeout("render request exceeded the router's 30 s limit")
        if b"timeout" in head or b"timed out" in head or b"deadline" in head:
            return RenderTimeout(f"render service timed out (limit {timeout}s)")
        return RenderBusy("render service is busy (HTTP 503)")
    if status == 429:
        return RenderBusy("render service is rate limiting (HTTP 429)")
    if status in (401, 403):
        return RenderUnavailable(f"render service rejected our credentials (HTTP {status})")
    if status == 502:
        return RenderUnavailable("render service is not reachable behind its router (HTTP 502)")
    return RenderRejected(f"render service could not convert the file (HTTP {status})")


def client_from_settings() -> GotenbergClient | None:
    url = service_url()
    if not url:
        return None
    return GotenbergClient(
        url,
        user=getattr(settings, "DOCUMENT_RENDER_SERVICE_USER", "") or "",
        password=getattr(settings, "DOCUMENT_RENDER_SERVICE_PASSWORD", "") or "",
        timeout=_int_setting("DOCUMENT_RENDER_HTTP_TIMEOUT", 60),
    )


# --------------------------------------------------------------------------- rasterize


def rasterize_pdf(pdf_bytes: bytes, expected_pages: int) -> list[bytes]:
    """One JPEG per page at the vision cap; the page count must match the batch."""
    import pypdfium2 as pdfium

    from chat.pdf_attach import PDF_RENDER_MAX_DIMENSION, render_pdf_pages_to_jpegs

    try:
        doc = pdfium.PdfDocument(pdf_bytes)
    except Exception as exc:  # noqa: BLE001
        raise RenderRejected("render service returned an unreadable PDF") from exc
    try:
        total = len(doc)
        if total != expected_pages:
            raise RenderRejected(f"expected {expected_pages} page(s) in the rendered PDF, got {total}")
        page = doc[0]
        try:
            width_pt, height_pt = page.get_size()
        finally:
            page.close()
    finally:
        doc.close()

    long_edge_pt = max(float(width_pt or 0), float(height_pt or 0)) or 960.0
    # Render slightly above the cap; render_pdf_pages_to_jpegs thumbnails down to
    # exactly PDF_RENDER_MAX_DIMENSION on the long edge.
    scale = max(0.5, (PDF_RENDER_MAX_DIMENSION / long_edge_pt) * 1.02)
    quality = _int_setting("VISION_IMAGE_JPEG_QUALITY", 82)
    jpegs, _total = render_pdf_pages_to_jpegs(
        pdf_bytes, max_pages=expected_pages, scale=scale, quality=quality,
    )
    if len(jpegs) != expected_pages:
        raise RenderRejected(f"rasterised {len(jpegs)} of {expected_pages} page(s)")
    return jpegs


# --------------------------------------------------------------------------- storage


def existing_page_numbers(version) -> set[int]:
    from chat.models import Asset

    return set(
        Asset.objects.filter(version=version, role=Asset.ROLE_PAGE_RENDER)
        .exclude(page_number__isnull=True)
        .values_list("page_number", flat=True)
    )


def store_page_render(version, page_number: int, jpeg: bytes):
    """Persist one slide image as a version-owned ``page_render`` Asset (idempotent)."""
    from django.core.files.base import ContentFile

    from chat.models import Asset

    existing = Asset.objects.filter(
        version=version, role=Asset.ROLE_PAGE_RENDER, page_number=page_number,
    ).first()
    if existing is not None:
        return existing

    width = height = None
    try:
        from PIL import Image

        with Image.open(io.BytesIO(jpeg)) as img:
            width, height = img.size
    except Exception:  # noqa: BLE001 — dimensions are informational
        pass

    asset = Asset(
        version=version,
        kind=Asset.KIND_IMAGE,
        role=Asset.ROLE_PAGE_RENDER,
        page_number=page_number,
        content_type="image/jpeg",
        size_bytes=len(jpeg),
        width=width,
        height=height,
        sha256=hashlib.sha256(jpeg).hexdigest(),
        description=f"Slide {page_number}",
    )
    try:
        asset.blob.save(f"{asset.id}.jpg", ContentFile(jpeg), save=True)
    except IntegrityError:
        # Lost a race with another worker rendering the same version: keep theirs
        # and drop the file we just wrote (the row insert failed after the upload).
        try:
            asset.blob.storage.delete(asset.blob.name)
        except Exception:  # noqa: BLE001
            logger.debug("page_render: could not remove the orphaned blob after a race", exc_info=True)
        return Asset.objects.get(version=version, role=Asset.ROLE_PAGE_RENDER, page_number=page_number)
    return asset


# --------------------------------------------------------------------------- state


def should_render(version) -> bool:
    """Whether this version qualifies: feature on, pptx, released and clean."""
    return (
        is_enabled()
        and (version.parser_type or "") == PPTX_PARSER_TYPE
        and version.status == DataRoomDocument.Status.READY
        and not version.is_quarantined
    )


def enqueue_page_render(version_id: int) -> bool:
    """Mark the version PENDING and queue the render task (best-effort dispatch).

    A publish failure leaves the version PENDING; the stale sweeper re-dispatches
    it later (bounded by ``page_render_attempts``).
    """
    from documents.tasks import render_document_pages

    DataRoomDocumentVersion.objects.filter(pk=version_id).update(
        page_render_state=DataRoomDocumentVersion.PageRenderState.PENDING,
        page_render_error="",
        updated_at=timezone.now(),
    )
    try:
        render_document_pages.delay(version_id)
    except Exception:  # noqa: BLE001 — broker blip; the sweeper retries
        logger.warning(
            "page_render: dispatch failed for version_id=%s; the stale sweeper will retry",
            version_id, exc_info=True,
        )
        return False
    return True


def _set_state(version_id: int, state: str, *, error: str | None = None, page_count: int | None = None) -> None:
    fields = {"page_render_state": state, "updated_at": timezone.now()}
    if error is not None:
        fields["page_render_error"] = error
    if page_count is not None:
        fields["page_count"] = page_count
    DataRoomDocumentVersion.objects.filter(pk=version_id).update(**fields)


def _touch_version(version_id: int) -> None:
    try:
        DataRoomDocumentVersion.objects.filter(pk=version_id).update(updated_at=timezone.now())
    except Exception:  # noqa: BLE001 — a failed heartbeat must not abort the render
        logger.debug("page_render: heartbeat failed for version_id=%s", version_id, exc_info=True)


@contextmanager
def _render_slot(version_id: int):
    """Hold the single per-process conversion slot; heartbeat while queued."""
    waited = 0.0
    while not _render_semaphore.acquire(timeout=_SLOT_WAIT_HEARTBEAT_S):
        waited += _SLOT_WAIT_HEARTBEAT_S
        _touch_version(version_id)
        logger.info("page_render: version_id=%s waiting for the render slot (%.0fs)", version_id, waited)
    try:
        yield
    finally:
        _render_semaphore.release()


def _read_native_bytes(version) -> bytes | None:
    source = version.native_blob if version.native_blob else version.document.original_file
    if not source:
        return None
    try:
        with source.open("rb") as fh:
            return fh.read()
    except Exception:  # noqa: BLE001
        logger.warning("page_render: could not read native bytes for version_id=%s", version.pk, exc_info=True)
        return None


# --------------------------------------------------------------------------- orchestrator


def _convert_slides(client: GotenbergClient, data: bytes, slide_numbers: list[int], version_id: int, upload_stem: str) -> list[bytes]:
    """Subset → convert (busy retried in place) → rasterize. Raises RenderError."""
    subset = split_pptx(data, slide_numbers)
    for attempt, delay in enumerate((*_BUSY_RETRY_DELAYS_S, None)):
        try:
            pdf = client.convert(subset, upload_name=f"{upload_stem}.pptx", trace=str(version_id))
            break
        except RenderBusy:
            if delay is None:
                raise
            logger.info(
                "page_render: version_id=%s render service busy; retrying in %ss (attempt %s)",
                version_id, delay, attempt + 1,
            )
            time.sleep(delay)
            _touch_version(version_id)
    return rasterize_pdf(pdf, len(slide_numbers))


def render_version_pages(version_id: int, *, count_attempt: bool = True) -> str:
    """Render every missing slide of a pptx version; returns the final state.

    Raises ``RenderBusy`` / ``RenderUnavailable`` (after in-place busy retries) so
    the Celery task can retry with backoff — the version then stays PENDING and
    resumes at the first missing slide.
    """
    from celery.exceptions import SoftTimeLimitExceeded

    State = DataRoomDocumentVersion.PageRenderState
    started = time.monotonic()
    try:
        version = DataRoomDocumentVersion.objects.select_related("document").get(pk=version_id)
    except DataRoomDocumentVersion.DoesNotExist:
        logger.info("page_render: version_id=%s not found (deleted before render)", version_id)
        return State.NONE

    if count_attempt:
        DataRoomDocumentVersion.objects.filter(pk=version_id).update(
            page_render_attempts=F("page_render_attempts") + 1,
        )
    if not is_enabled():
        _set_state(version_id, State.FAILED, error="Document render service is not configured.")
        return State.FAILED
    if not should_render(version):
        _set_state(version_id, State.SKIPPED, error="Not a renderable presentation version.")
        return State.SKIPPED

    raw = _read_native_bytes(version)
    if not raw:
        _set_state(version_id, State.FAILED, error="The original file is unavailable.")
        return State.FAILED
    try:
        total = count_slides(raw)
    except Exception:  # noqa: BLE001
        logger.warning("page_render: version_id=%s could not open the presentation", version_id, exc_info=True)
        _set_state(version_id, State.FAILED, error="The presentation could not be opened.")
        return State.FAILED
    if total <= 0:
        _set_state(version_id, State.FAILED, page_count=0, error="The presentation has no slides.")
        return State.FAILED
    max_slides = _int_setting("DOCUMENT_RENDER_MAX_SLIDES", 200)
    if total > max_slides:
        _set_state(
            version_id, State.SKIPPED, page_count=total,
            error=f"The deck has {total} slides; slide previews are only rendered for decks up to {max_slides}.",
        )
        return State.SKIPPED
    _set_state(version_id, State.PENDING, page_count=total)

    already = existing_page_numbers(version)
    todo = [n for n in range(1, total + 1) if n not in already]
    if not todo:
        _set_state(version_id, State.READY, error="")
        return State.READY

    data = optimize_pptx(raw)
    del raw
    client = client_from_settings()
    batch_size = _int_setting("DOCUMENT_RENDER_BATCH_SIZE", 8)
    batches = [todo[i:i + batch_size] for i in range(0, len(todo), batch_size)]
    rendered: list[int] = []
    failed: list[int] = []

    try:
        with _render_slot(version_id):
            for batch_no, batch in enumerate(batches, start=1):
                batch_started = time.monotonic()
                try:
                    jpegs = _convert_slides(client, data, batch, version_id, f"{version_id}-{batch_no}")
                except (RenderRejected, RenderTimeout) as exc:
                    logger.warning(
                        "page_render: version_id=%s batch %s/%s (%s slides) failed with %s; retrying slide by slide",
                        version_id, batch_no, len(batches), len(batch), exc.__class__.__name__,
                    )
                    jpegs = None
                if jpegs is not None:
                    for slide_no, jpeg in zip(batch, jpegs):
                        store_page_render(version, slide_no, jpeg)
                        rendered.append(slide_no)
                else:
                    for slide_no in batch:
                        try:
                            (jpeg,) = _convert_slides(
                                client, data, [slide_no], version_id, f"{version_id}-{batch_no}-{slide_no}",
                            )
                        except (RenderRejected, RenderTimeout) as exc:
                            logger.warning(
                                "page_render: version_id=%s slide %s could not be rendered (%s)",
                                version_id, slide_no, exc.__class__.__name__,
                            )
                            failed.append(slide_no)
                            continue
                        store_page_render(version, slide_no, jpeg)
                        rendered.append(slide_no)
                _touch_version(version_id)
                logger.info(
                    "page_render: version_id=%s batch %s/%s slides=%s done in %.1fs",
                    version_id, batch_no, len(batches), len(batch), time.monotonic() - batch_started,
                )
    except SoftTimeLimitExceeded:
        remaining = [n for n in todo if n not in rendered and n not in failed]
        logger.warning(
            "page_render: version_id=%s hit the soft time limit with %s slide(s) left; finalising",
            version_id, len(remaining),
        )
        failed.extend(remaining)
    except (RenderBusy, RenderUnavailable) as exc:
        logger.warning(
            "page_render: version_id=%s paused after %s new slide(s): %s; Celery will retry",
            version_id, len(rendered), exc.__class__.__name__,
        )
        raise

    have = already | set(rendered)
    if failed and have:
        state = State.PARTIAL
        error = "Slides " + ", ".join(str(n) for n in sorted(failed)) + " could not be rendered."
    elif failed:
        state = State.FAILED
        error = "No slides could be rendered."
    else:
        state = State.READY
        error = ""
    _set_state(version_id, state, error=error)
    logger.info(
        "page_render: version_id=%s state=%s slides=%s batches=%s rendered=%s failed=%s in %.1fs",
        version_id, state, total, len(batches), len(rendered), len(failed), time.monotonic() - started,
    )
    return state


__all__ = [
    "GotenbergClient",
    "RenderBusy",
    "RenderError",
    "RenderRejected",
    "RenderTimeout",
    "RenderUnavailable",
    "client_from_settings",
    "count_slides",
    "enqueue_page_render",
    "existing_page_numbers",
    "is_enabled",
    "literalize_slide_numbers",
    "optimize_pptx",
    "rasterize_pdf",
    "render_version_pages",
    "should_render",
    "split_pptx",
    "store_page_render",
]
