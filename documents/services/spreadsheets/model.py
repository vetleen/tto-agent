"""Compact in-memory model of a workbook, built without the openpyxl DOM.

Two complementary reads of the same zip:

* **openpyxl** ``read_only=True, data_only=True`` streams cell *values* and
  resolved number formats. Read-only worksheets expose nothing structural
  (no ``merged_cells``, ``column_dimensions``, ``row_dimensions``, ``tables``,
  ``auto_filter``, ``freeze_panes``, ``conditional_formatting`` — verified
  against 3.1.5), and the full DOM costs ~30-50x file size in RAM, so:
* a **targeted XML pass** (``iterparse`` straight from the zip) supplies the
  structure: merged ranges, column widths/hidden, row heights/hidden, freeze
  pane, autofilter, table refs + column names, and ``cellIs`` conditional-
  formatting rules with their dxf fill colours.

The model keeps only *visible* cells as formatted text in plain tuples —
original Excel coordinates preserved throughout — and is dropped as soon as
chunks/manifest are built.
"""
from __future__ import annotations

import logging
import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from xml.etree.ElementTree import iterparse

logger = logging.getLogger(__name__)

# Conditional-formatting matches recorded per sheet before we stop bothering —
# a rule matching every cell of a huge sheet must not balloon the model.
_CF_MATCH_CAP = 20_000
# Leading rows whose value-type mix is recorded for header detection.
_TYPE_PROFILE_ROWS = 10

_EXCEL_EPOCH_FORMATS = ("General", "")


def col_letter(index: int) -> str:
    """1-based column index -> Excel letter ("A", "AB")."""
    out = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def col_index(letters: str) -> int:
    """Excel column letters -> 1-based index ("A" -> 1)."""
    out = 0
    for ch in letters.strip().upper():
        if not "A" <= ch <= "Z":
            raise ValueError(f"Invalid column letters: {letters!r}")
        out = out * 26 + (ord(ch) - 64)
    return out


_CELL_REF_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")


def parse_cell_ref(ref: str) -> tuple[int, int]:
    """"B7" -> (row=7, col=2). Raises ValueError on anything else."""
    m = _CELL_REF_RE.match(ref.strip())
    if not m:
        raise ValueError(f"Invalid cell reference: {ref!r}")
    return int(m.group(2)), col_index(m.group(1))


def parse_range(ref: str) -> tuple[int, int, int, int]:
    """"A1:C4" (or a single cell) -> (min_row, min_col, max_row, max_col)."""
    first, _, last = ref.strip().partition(":")
    r1, c1 = parse_cell_ref(first)
    r2, c2 = parse_cell_ref(last) if last else (r1, c1)
    return min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)


# ---------------------------------------------------------------------------
# Dataclasses
# ---------------------------------------------------------------------------

@dataclass(slots=True)
class TableInfo:
    name: str
    bounds: tuple[int, int, int, int]  # min_row, min_col, max_row, max_col
    header_row_count: int
    column_names: list[str]


@dataclass(slots=True)
class CfRule:
    ranges: list[tuple[int, int, int, int]]
    operator: str
    operands: list  # str | float constants
    color: str  # "#RRGGBB"


@dataclass(slots=True)
class SheetModel:
    name: str
    index: int  # 0-based position among the processed (visible) sheets
    # Sparse visible content, original Excel coordinates:
    # [(excel_row, [(excel_col, formatted_text), ...]), ...] ascending.
    rows: list = field(default_factory=list)
    max_data_row: int = 0
    max_data_col: int = 0
    cell_count: int = 0
    # Structure from the XML pass:
    col_widths: dict = field(default_factory=dict)  # col -> Excel width units
    hidden_cols: set = field(default_factory=set)
    row_heights: dict = field(default_factory=dict)  # row -> points (custom only)
    hidden_rows: set = field(default_factory=set)
    merged: list = field(default_factory=list)  # bounds tuples
    freeze_rows: int = 0
    freeze_cols: int = 0
    auto_filter: tuple | None = None  # bounds
    tables: list = field(default_factory=list)  # TableInfo
    cf_fills: dict = field(default_factory=dict)  # (row, col) -> "#RRGGBB"
    default_row_height: float = 15.0
    # Header detection input: [(excel_row, n_nonempty, n_text)] for the first
    # visible non-empty rows.
    type_profile: list = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not self.rows

    def visible_cols(self) -> list[int]:
        return [c for c in range(1, self.max_data_col + 1) if c not in self.hidden_cols]

    def row_cells(self, excel_row: int) -> list:
        """[(col, text)] for one row (empty list when absent)."""
        for r, cells in self.rows:
            if r == excel_row:
                return cells
            if r > excel_row:
                break
        return []


@dataclass(slots=True)
class WorkbookModel:
    sheets: list  # SheetModel, processed (visible) sheets in workbook order
    skipped_sheets: list  # names beyond XLSX_MAX_SHEETS
    hidden_sheets: list  # names of hidden/veryHidden sheets (never processed)
    total_cells: int = 0


# ---------------------------------------------------------------------------
# Cell formatting — deliberately small. Non-goal: the full Excel format
# language (multi-section negatives, locale prefixes, fractions, scientific).
# ---------------------------------------------------------------------------

_QUOTED_RE = re.compile(r'"([^"]*)"')
_BRACKET_RE = re.compile(r"\[[^\]]*\]")
_CURRENCY_CHARS = "$€£¥"


@lru_cache(maxsize=512)
def _parse_number_format(number_format: str) -> tuple:
    """First section of an Excel number format -> (kind, decimals, thousands,
    prefix, suffix). kind: "general" | "percent" | "number" | "text"."""
    section = (number_format or "General").split(";")[0]
    if section.strip() in _EXCEL_EPOCH_FORMATS or section.strip().lower() == "general":
        return ("general", 0, False, "", "")
    if "@" in section:
        return ("text", 0, False, "", "")

    prefix = suffix = ""
    digits_seen = False
    first_digit = _first_digit_pos(_QUOTED_RE.sub(lambda m: " " * len(m.group(0)), section))
    literals = [(m.start() < first_digit, m.group(1)) for m in _QUOTED_RE.finditer(section)]
    stripped = _QUOTED_RE.sub("", section)
    stripped = _BRACKET_RE.sub("", stripped)
    for before, text in literals:
        if before:
            prefix += text
        else:
            suffix += text
    # Bare currency characters count as literals too.
    for ch in stripped:
        if ch in _CURRENCY_CHARS:
            if digits_seen:
                suffix += ch
            else:
                prefix += ch
        elif ch in "0#?":
            digits_seen = True

    body = "".join(ch for ch in stripped if ch not in _CURRENCY_CHARS)
    has_placeholder = any(ch in "0#?" for ch in body)
    if "e+" in body.lower() or "e-" in body.lower():
        return ("general", 0, False, "", "")  # scientific: fall back
    if not has_placeholder:
        # A pure date/time picture on a still-numeric value (openpyxl usually
        # converts these to datetime already), or nothing to format: fall back.
        return ("general", 0, False, "", "")
    is_percent = "%" in body
    int_part, _, frac_part = body.partition(".")
    decimals = sum(1 for ch in frac_part if ch in "0#")
    thousands = "," in int_part
    kind = "percent" if is_percent else "number"
    return (kind, decimals, thousands, prefix, suffix)


def _first_digit_pos(section: str) -> int:
    for i, ch in enumerate(section):
        if ch in "0#?":
            return i
    return len(section)


def _format_general_number(value) -> str:
    if isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
        return str(int(value))
    return f"{value:.10g}"


def format_cell_value(value, number_format: str = "General") -> str:
    """Human-readable text for a cell value, honouring the common formats:
    General, fixed decimals, thousands separators, percent, literal currency,
    ``@`` text, and date/time (openpyxl already yields datetime objects for
    date-formatted cells). Everything exotic falls back to a plain number.
    """
    import datetime

    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return value.replace("\x00", "")
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time.min:
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M")
    if isinstance(value, datetime.timedelta):
        return str(value)
    if not isinstance(value, (int, float)):
        return str(value)

    kind, decimals, thousands, prefix, suffix = _parse_number_format(number_format)
    if kind == "percent":
        return f"{value * 100:.{decimals}f}%"
    if kind in ("general", "text"):
        return _format_general_number(value)
    sep = "," if thousands else ""
    return f"{prefix}{value:{sep}.{decimals}f}{suffix}"


# ---------------------------------------------------------------------------
# XML supplement — everything read_only mode does not expose.
# ---------------------------------------------------------------------------

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _attr(elem, name: str, default=None):
    """Namespace-tolerant attribute get (handles r:id)."""
    if name in elem.attrib:
        return elem.attrib[name]
    for key, val in elem.attrib.items():
        if _local(key) == name:
            return val
    return default


@dataclass(slots=True)
class _SheetXml:
    # (min_col, max_col, width|None, hidden) — ranges as written in <cols>.
    col_ranges: list = field(default_factory=list)
    row_heights: dict = field(default_factory=dict)
    hidden_rows: set = field(default_factory=set)
    merged: list = field(default_factory=list)
    freeze_rows: int = 0
    freeze_cols: int = 0
    auto_filter: tuple | None = None
    table_rel_ids: list = field(default_factory=list)
    cf_rules: list = field(default_factory=list)  # CfRule (color filled later)
    default_row_height: float = 15.0


def _normalize_part(target: str, base_dir: str = "xl") -> str:
    """Resolve a relationship Target ("/xl/worksheets/sheet1.xml",
    "worksheets/sheet1.xml", "../tables/table1.xml") to a zip member path."""
    target = target.strip()
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(base_dir, target))


def _parse_rels(zf: zipfile.ZipFile, part: str) -> dict:
    """Relationship Id -> Target for a part's .rels (empty when absent)."""
    rels_path = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    try:
        data = zf.open(rels_path)
    except KeyError:
        return {}
    out = {}
    with data:
        for _, elem in iterparse(data):
            if _local(elem.tag) == "Relationship":
                out[elem.get("Id", "")] = elem.get("Target", "")
    return out


def _sheet_parts(zf: zipfile.ZipFile) -> dict:
    """Sheet name -> zip path of its worksheet XML, via workbook.xml + rels."""
    rels = _parse_rels(zf, "xl/workbook.xml")
    out = {}
    with zf.open("xl/workbook.xml") as fh:
        for _, elem in iterparse(fh):
            if _local(elem.tag) == "sheet":
                rid = _attr(elem, "id", "")
                target = rels.get(rid)
                if target:
                    out[elem.get("name", "")] = _normalize_part(target)
    return out


def _parse_cf_operand(text: str):
    """A cfRule <formula> constant -> str/float, or None when not a constant."""
    text = (text or "").strip()
    if len(text) >= 2 and text.startswith('"') and text.endswith('"'):
        return text[1:-1].replace('""', '"')
    try:
        return float(text)
    except ValueError:
        return None


_CELLIS_OPERATORS = frozenset({
    "equal", "notEqual", "greaterThan", "greaterThanOrEqual",
    "lessThan", "lessThanOrEqual", "between", "notBetween",
})


def _parse_sheet_xml(fh) -> _SheetXml:
    """One streaming pass over a worksheet part: structure attributes only
    (cell values come from openpyxl). Row elements are cleared as they close
    so sheetData never accumulates in memory.
    """
    info = _SheetXml()
    cf_sqref: str | None = None
    for event, elem in iterparse(fh, events=("start", "end")):
        tag = _local(elem.tag)
        if event == "start":
            if tag == "row":
                r = elem.get("r")
                if r is not None:
                    r = int(r)
                    if elem.get("hidden") in ("1", "true"):
                        info.hidden_rows.add(r)
                    ht = elem.get("ht")
                    if ht is not None and elem.get("customHeight") in ("1", "true"):
                        info.row_heights[r] = float(ht)
            elif tag == "conditionalFormatting":
                cf_sqref = elem.get("sqref", "")
            continue

        # end events
        if tag == "row" or tag == "sheetData":
            elem.clear()
        elif tag == "col":
            lo = int(elem.get("min", "0") or 0)
            hi = int(elem.get("max", "0") or 0)
            width = elem.get("width")
            hidden = elem.get("hidden") in ("1", "true")
            # A width attribute means the column has a set width (custom or
            # autofit) — both are intent signals. Ranges can span "all columns"
            # (max=16384); expansion is clamped to the sheet's real width later.
            if hi >= lo > 0:
                info.col_ranges.append(
                    (lo, hi, float(width) if width is not None else None, hidden)
                )
        elif tag == "pane":
            if elem.get("state") in ("frozen", "frozenSplit"):
                info.freeze_cols = int(float(elem.get("xSplit", "0") or 0))
                info.freeze_rows = int(float(elem.get("ySplit", "0") or 0))
        elif tag == "autoFilter":
            # Only the worksheet-level element reaches this parser (a table's
            # autoFilter lives in its own tables/*.xml part).
            ref = elem.get("ref")
            if ref and info.auto_filter is None:
                try:
                    info.auto_filter = parse_range(ref)
                except ValueError:
                    pass
        elif tag == "mergeCell":
            ref = elem.get("ref")
            if ref:
                try:
                    info.merged.append(parse_range(ref))
                except ValueError:
                    pass
        elif tag == "cfRule":
            if elem.get("type") == "cellIs" and cf_sqref:
                op = elem.get("operator", "")
                dxf = elem.get("dxfId")
                operands = [
                    _parse_cf_operand(f.text)
                    for f in elem
                    if _local(f.tag) == "formula"
                ]
                if op in _CELLIS_OPERATORS and dxf is not None and operands and all(
                    o is not None for o in operands
                ):
                    try:
                        ranges = [parse_range(part) for part in cf_sqref.split()]
                    except ValueError:
                        ranges = []
                    if ranges:
                        # color resolved from the dxf index by the caller
                        info.cf_rules.append((ranges, op, operands, int(dxf)))
                else:
                    logger.debug("spreadsheets: skipping unsupported cellIs rule (op=%s)", op)
            elif elem.get("type") not in (None, "cellIs"):
                logger.debug("spreadsheets: skipping CF rule type=%s", elem.get("type"))
        elif tag == "conditionalFormatting":
            cf_sqref = None
            elem.clear()
        elif tag == "tablePart":
            rid = _attr(elem, "id")
            if rid:
                info.table_rel_ids.append(rid)
        elif tag == "sheetFormatPr":
            drh = elem.get("defaultRowHeight")
            if drh:
                try:
                    info.default_row_height = float(drh)
                except ValueError:
                    pass
    return info


def _parse_table_xml(fh) -> TableInfo | None:
    name = ref = None
    header_rows = 1
    columns: list[str] = []
    for _, elem in iterparse(fh):
        tag = _local(elem.tag)
        if tag == "table":
            name = elem.get("displayName") or elem.get("name") or ""
            ref = elem.get("ref")
            hrc = elem.get("headerRowCount")
            if hrc is not None:
                header_rows = int(hrc)
        elif tag == "tableColumn":
            columns.append(elem.get("name", ""))
    if not ref:
        return None
    try:
        bounds = parse_range(ref)
    except ValueError:
        return None
    return TableInfo(name=name or "", bounds=bounds, header_row_count=header_rows, column_names=columns)


def _parse_dxf_fills(zf: zipfile.ZipFile) -> list:
    """styles.xml <dxfs> in order -> fill colour per dxfId ("#RRGGBB" or None).

    Only solid pattern fills with an explicit rgb are supported; theme/indexed
    colours yield None and their rules are skipped.
    """
    try:
        fh = zf.open("xl/styles.xml")
    except KeyError:
        return []
    fills: list = []
    in_dxfs = False
    current: str | None = None
    pattern_ok = False
    with fh:
        for event, elem in iterparse(fh, events=("start", "end")):
            tag = _local(elem.tag)
            if event == "start":
                if tag == "dxfs":
                    in_dxfs = True
                elif in_dxfs and tag == "dxf":
                    current = None
                    pattern_ok = False
                elif in_dxfs and tag == "patternFill":
                    pattern_ok = elem.get("patternType") in (None, "solid")
                elif in_dxfs and pattern_ok and tag in ("fgColor", "bgColor"):
                    rgb = elem.get("rgb")
                    if rgb and len(rgb) >= 6 and (current is None or tag == "fgColor"):
                        current = "#" + rgb[-6:].upper()
                continue
            if tag == "dxf" and in_dxfs:
                fills.append(current)
            elif tag == "dxfs":
                in_dxfs = False
                elem.clear()
    return fills


def _cf_matches(operator: str, operands: list, value) -> bool:
    """Evaluate a cellIs rule against a raw cell value."""
    if isinstance(value, bool) or value is None:
        return False
    if isinstance(value, str):
        if operator not in ("equal", "notEqual") or not operands or not isinstance(operands[0], str):
            return False
        same = value.strip().casefold() == operands[0].strip().casefold()
        return same if operator == "equal" else not same
    if not isinstance(value, (int, float)):
        return False
    nums = [o for o in operands if isinstance(o, (int, float))]
    if len(nums) != len(operands):
        return False
    if operator == "equal":
        return value == nums[0]
    if operator == "notEqual":
        return value != nums[0]
    if operator == "greaterThan":
        return value > nums[0]
    if operator == "greaterThanOrEqual":
        return value >= nums[0]
    if operator == "lessThan":
        return value < nums[0]
    if operator == "lessThanOrEqual":
        return value <= nums[0]
    if len(nums) < 2:
        return False
    lo, hi = min(nums[0], nums[1]), max(nums[0], nums[1])
    if operator == "between":
        return lo <= value <= hi
    if operator == "notBetween":
        return not (lo <= value <= hi)
    return False


# ---------------------------------------------------------------------------
# Loader
# ---------------------------------------------------------------------------

def _check_zip_budget(path) -> None:
    from django.conf import settings

    max_bytes = getattr(settings, "XLSX_MAX_UNCOMPRESSED_BYTES", 250_000_000)
    try:
        with zipfile.ZipFile(path) as zf:
            total = sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile as exc:
        raise ValueError("This file is corrupt or not a valid Excel workbook.") from exc
    if total > max_bytes:
        raise ValueError(
            "This workbook is too large to process "
            f"(uncompressed content over {max_bytes // 1_000_000} MB)."
        )


def _apply_col_ranges(xml: _SheetXml, max_col: int) -> tuple[dict, set]:
    """Expand <col min max> ranges into per-column width/hidden maps, clamped
    to the sheet's real width (ranges routinely span to column 16384)."""
    widths: dict = {}
    hidden: set = set()
    for lo, hi, width, is_hidden in xml.col_ranges:
        hi = min(hi, max_col)
        for c in range(lo, hi + 1):
            if width is not None:
                widths[c] = width
            if is_hidden:
                hidden.add(c)
    return widths, hidden


def load_workbook_model(path, only_sheets=None) -> WorkbookModel:
    """Build the compact model for a workbook file.

    ``only_sheets`` (a set of sheet names) streams just those visible sheets —
    used by chat-time lazy tile renders — while keeping each sheet's ``index``
    equal to its position among ALL visible sheets, matching the manifest.

    Raises ``ValueError`` (user-visible processing error) on corrupt files,
    zip bombs, and cell budgets — the pipeline's existing handler surfaces it.
    """
    from django.conf import settings

    from openpyxl import load_workbook

    max_cells = getattr(settings, "XLSX_MAX_CELLS", 1_000_000)
    max_sheets = getattr(settings, "XLSX_MAX_SHEETS", 50)
    max_cell_chars = getattr(settings, "XLSX_MAX_CELL_CHARS", 500)
    max_scan_cells = getattr(settings, "XLSX_MAX_SCAN_CELLS", 50_000_000)
    max_scan_cols = getattr(settings, "XLSX_MAX_SCAN_COLS", 1024)

    _check_zip_budget(path)

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except zipfile.BadZipFile as exc:
        raise ValueError("This file is corrupt or not a valid Excel workbook.") from exc
    except Exception as exc:
        raise ValueError(f"Could not open this Excel workbook: {exc}") from exc

    sheets: list[SheetModel] = []
    skipped: list[str] = []
    hidden_sheets: list[str] = []
    total_cells = 0
    try:
        visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
        hidden_sheets = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
        skipped = [ws.title for ws in visible[max_sheets:]]
        visible = visible[:max_sheets]

        with zipfile.ZipFile(path) as zf:
            parts = _sheet_parts(zf)
            dxf_fills = _parse_dxf_fills(zf)

            for idx, ws in enumerate(visible):
                if only_sheets is not None and ws.title not in only_sheets:
                    continue
                sheet = SheetModel(name=ws.title, index=idx)
                part = parts.get(ws.title)
                xml = _SheetXml()
                if part and part in zf.namelist():
                    try:
                        with zf.open(part) as fh:
                            xml = _parse_sheet_xml(fh)
                    except Exception:
                        logger.exception(
                            "spreadsheets: XML supplement failed for sheet %r; continuing with values only",
                            ws.title,
                        )
                sheet.row_heights = xml.row_heights
                sheet.hidden_rows = xml.hidden_rows
                sheet.merged = xml.merged
                sheet.freeze_rows = xml.freeze_rows
                sheet.freeze_cols = xml.freeze_cols
                sheet.auto_filter = xml.auto_filter
                sheet.default_row_height = xml.default_row_height

                # Tables (header labels) via the sheet's rels.
                if part and xml.table_rel_ids:
                    rels = _parse_rels(zf, part)
                    for rid in xml.table_rel_ids:
                        target = rels.get(rid)
                        if not target:
                            continue
                        tpath = _normalize_part(target, posixpath.dirname(part))
                        if tpath in zf.namelist():
                            with zf.open(tpath) as fh:
                                table = _parse_table_xml(fh)
                            if table:
                                sheet.tables.append(table)

                cf_rules = [
                    CfRule(ranges=ranges, operator=op, operands=operands, color=dxf_fills[dxf])
                    for ranges, op, operands, dxf in xml.cf_rules
                    if 0 <= dxf < len(dxf_fills) and dxf_fills[dxf]
                ]

                # Stream values. Hidden rows/cols are dropped (chunks, mesh and
                # tiles all show what the user sees); original coordinates kept.
                hidden_rows = sheet.hidden_rows
                cf_capped = False
                # Bound the scan box. openpyxl read-only iter_rows() pads to the
                # declared <dimension>, so a stray far-corner cell would spin over
                # billions of empty cells. Cap columns and derive a row cap so the
                # padded scan stays within max_scan_cells; genuine bulk data still
                # trips XLSX_MAX_CELLS below.
                dim_rows = ws.max_row or 0
                dim_cols = ws.max_column or 0
                scan_cols = max(1, min(dim_cols, max_scan_cols) if dim_cols else max_scan_cols)
                row_cap = max(1, max_scan_cells // scan_cols)
                scan_rows = min(dim_rows, row_cap) if dim_rows else row_cap
                if dim_rows > scan_rows or dim_cols > scan_cols:
                    logger.warning(
                        "spreadsheets: sheet %r declared range %d×%d exceeds scan "
                        "box %d×%d; clamping (possible stray far cell)",
                        ws.title, dim_rows, dim_cols, scan_rows, scan_cols,
                    )
                for row in ws.iter_rows(max_row=scan_rows, max_col=scan_cols):
                    row_cells: list = []
                    excel_row = None
                    n_text = 0
                    for cell in row:
                        value = cell.value
                        if value is None or value == "":
                            continue
                        excel_row = cell.row
                        col = cell.column
                        if col > sheet.max_data_col:
                            sheet.max_data_col = col
                        if isinstance(value, str):
                            n_text += 1
                        if cf_rules and len(sheet.cf_fills) < _CF_MATCH_CAP:
                            for rule in cf_rules:
                                if any(
                                    r1 <= cell.row <= r2 and c1 <= col <= c2
                                    for r1, c1, r2, c2 in rule.ranges
                                ) and _cf_matches(rule.operator, rule.operands, value):
                                    sheet.cf_fills[(cell.row, col)] = rule.color
                                    break
                        elif cf_rules and not cf_capped:
                            cf_capped = True
                            logger.debug(
                                "spreadsheets: CF match cap reached on sheet %r", ws.title
                            )
                        text = format_cell_value(value, cell.number_format)
                        if len(text) > max_cell_chars:
                            text = text[: max_cell_chars - 1] + "…"
                        row_cells.append((col, text))
                    if excel_row is None:
                        continue
                    if excel_row in hidden_rows:
                        continue
                    if len(sheet.type_profile) < _TYPE_PROFILE_ROWS:
                        sheet.type_profile.append((excel_row, len(row_cells), n_text))
                    total_cells += len(row_cells)
                    if total_cells > max_cells:
                        raise ValueError(
                            "This workbook is too large to process "
                            f"(over {max_cells:,} cells)."
                        )
                    sheet.max_data_row = excel_row
                    sheet.cell_count += len(row_cells)
                    sheet.rows.append((excel_row, row_cells))

                # Hidden columns' cells are dropped after the fact so the
                # width/hidden clamp can use the true max column.
                widths, hidden_cols = _apply_col_ranges(xml, sheet.max_data_col or 1)
                sheet.col_widths = widths
                sheet.hidden_cols = hidden_cols
                if hidden_cols:
                    for i, (excel_row, cells) in enumerate(sheet.rows):
                        kept = [(c, t) for c, t in cells if c not in hidden_cols]
                        if len(kept) != len(cells):
                            sheet.cell_count -= len(cells) - len(kept)
                            total_cells -= len(cells) - len(kept)
                            sheet.rows[i] = (excel_row, kept)
                    sheet.rows = [entry for entry in sheet.rows if entry[1]]
                sheets.append(sheet)
    finally:
        wb.close()

    return WorkbookModel(
        sheets=sheets,
        skipped_sheets=skipped,
        hidden_sheets=hidden_sheets,
        total_cells=total_cells,
    )
