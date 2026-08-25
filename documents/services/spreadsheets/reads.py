"""Exact, streaming reads for the chat tools: the Ctrl+F primitive and
row/cell reads. openpyxl ``read_only`` only — O(1) memory regardless of
workbook size, formatted through the same ``format_cell_value`` the chunks
use so the agent always sees the user's numbers.
"""
from __future__ import annotations

import logging

from documents.services.spreadsheets.model import col_index, format_cell_value

logger = logging.getLogger(__name__)


def parse_col_spec(spec: str) -> list[int]:
    """"A,C-F" -> [1, 3, 4, 5, 6]. Raises ValueError on nonsense."""
    out: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        first, _, last = part.partition("-")
        lo = col_index(first)
        hi = col_index(last) if last else lo
        if hi < lo:
            lo, hi = hi, lo
        out.extend(range(lo, hi + 1))
    if not out:
        raise ValueError(f"No columns in spec: {spec!r}")
    return sorted(set(out))


def _formatted(cell, max_chars: int) -> str:
    text = format_cell_value(cell.value, cell.number_format)
    if len(text) > max_chars:
        text = text[: max_chars - 1] + "…"
    return text


def _open_read_only(path):
    from openpyxl import load_workbook

    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not open this Excel workbook: {exc}") from exc


def find_in_workbook(path, needle: str, sheet_names=None, max_results: int = 50,
                     max_cell_chars: int = 500) -> tuple[list, bool]:
    """Case-insensitive substring scan over formatted cell text.

    Returns ``([{sheet, row, col, value}, ...], truncated)`` in sheet/row
    order. ``sheet_names`` (set) limits the scan; hidden sheets are never
    scanned (they are not processed either).
    """
    folded = needle.casefold()
    hits: list = []
    truncated = False
    wb = _open_read_only(path)
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            if sheet_names is not None and ws.title not in sheet_names:
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None or cell.value == "":
                        continue
                    text = _formatted(cell, max_cell_chars)
                    if folded in text.casefold():
                        hits.append({
                            "sheet": ws.title, "row": cell.row, "col": cell.column,
                            "value": text,
                        })
                        if len(hits) >= max_results:
                            truncated = True
                            return hits, truncated
    finally:
        wb.close()
    return hits, truncated


def read_rows(path, sheet_name: str, row_start: int, row_end: int, cols=None,
              max_rows: int = 200, max_cell_chars: int = 500) -> tuple[list, bool]:
    """Exact rows from one sheet: ``([(excel_row, [(col, text), ...]), ...],
    truncated)``. Empty rows in range are omitted; ``cols`` (list of indices)
    filters columns."""
    col_set = set(cols) if cols else None
    out: list = []
    truncated = False
    wb = _open_read_only(path)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"No sheet named {sheet_name!r} in this workbook.")
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=row_start, max_row=row_end):
            cells = []
            excel_row = None
            for cell in row:
                if cell.value is None or cell.value == "":
                    continue
                if col_set is not None and cell.column not in col_set:
                    continue
                excel_row = cell.row
                cells.append((cell.column, _formatted(cell, max_cell_chars)))
            if excel_row is None:
                continue
            if len(out) >= max_rows:
                truncated = True
                break
            out.append((excel_row, cells))
    finally:
        wb.close()
    return out, truncated


def read_cell(path, sheet_name: str, row: int, col: int,
              max_cell_chars: int = 500) -> tuple[str | None, list]:
    """One cell's formatted text (None when empty) plus its full row for
    context: ``(value, [(col, text), ...])``."""
    rows, _ = read_rows(path, sheet_name, row, row, max_rows=1,
                        max_cell_chars=max_cell_chars)
    if not rows:
        return None, []
    _, cells = rows[0]
    value = dict(cells).get(col)
    return value, cells
