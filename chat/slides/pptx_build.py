"""Deck JSON -> ``.pptx`` bytes (python-pptx).

``build_deck_pptx(deck, ...)`` is the pure core (no Django) so it unit-tests
with golden-OOXML assertions on any platform. ``build_pptx(slide_set)`` is the
Django-aware wrapper that pulls the deck + base template off a ``SlideSet`` and
binds an image resolver scoped to the deck's owner.

Coordinates are points (Pt); python-pptx converts to EMU. Colours resolve via
:mod:`chat.slides.theme` — native theme slots become ``MSO_THEME_COLOR`` refs
(tracking the template's scheme + showing in PowerPoint's picker), everything
else an explicit RGB.
"""

from __future__ import annotations

import logging
import re
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from chat.slides import oxml_pokes as pokes
from chat.slides import theme as theme_mod

logger = logging.getLogger(__name__)

ASSETS_DIR = Path(__file__).resolve().parent / "assets"
TEMPLATE_DIR = ASSETS_DIR / "templates"

# Curated shape vocabulary (friendly name -> MSO_SHAPE). Verified members only;
# unknown names fall back to RECTANGLE with a warning.
SHAPE_MAP = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "ellipse": MSO_SHAPE.OVAL,
    "circle": MSO_SHAPE.OVAL,
    "right_arrow": MSO_SHAPE.RIGHT_ARROW,
    "left_arrow": MSO_SHAPE.LEFT_ARROW,
    "up_arrow": MSO_SHAPE.UP_ARROW,
    "down_arrow": MSO_SHAPE.DOWN_ARROW,
    "chevron": MSO_SHAPE.CHEVRON,
    "pentagon": MSO_SHAPE.PENTAGON,
    "diamond": MSO_SHAPE.DIAMOND,
    "hexagon": MSO_SHAPE.HEXAGON,
    "star": MSO_SHAPE.STAR_5_POINT,
    "plus": MSO_SHAPE.MATH_PLUS,
    "cross": MSO_SHAPE.CROSS,
    "cloud": MSO_SHAPE.CLOUD,
    "heart": MSO_SHAPE.HEART,
    "donut": MSO_SHAPE.DONUT,
    "plaque": MSO_SHAPE.PLAQUE,
    "frame": MSO_SHAPE.FRAME,
    "bevel": MSO_SHAPE.BEVEL,
    "cube": MSO_SHAPE.CUBE,
    "can": MSO_SHAPE.CAN,
    "callout": MSO_SHAPE.LINE_CALLOUT_1,
    "block_arc": MSO_SHAPE.BLOCK_ARC,
    "tear": MSO_SHAPE.TEAR,
    "lightning": MSO_SHAPE.LIGHTNING_BOLT,
    "smiley": MSO_SHAPE.SMILEY_FACE,
}

# The deck theme names our shipped OFL font clones (Caladea/Carlito/…). Those
# aren't installed on most machines, so a downloaded .pptx that referenced them
# would be substituted by the viewer's PowerPoint — often with the *wrong* style
# (a serif headline turning sans). Reference the widely-installed Office fonts
# our clones are metric-compatible with instead, so the .pptx renders correctly
# for viewers while the Pillow preview keeps using the (identical-metric) clone
# faces it ships. Names not listed pass through unchanged.
_PPTX_FONT_ALIAS = {
    "Caladea": "Cambria",
    "Carlito": "Calibri",
    "Tinos": "Times New Roman",
    "Arimo": "Arial",
    "Cousine": "Courier New",
    "Gelasio": "Georgia",
    "EBGaramond": "Garamond",
}


def _pptx_font(theme, font_ref):
    fam = theme_mod.font_family(theme, font_ref)
    return _PPTX_FONT_ALIAS.get(fam, fam)


THEME_COLOR_MAP = {
    "dk1": MSO_THEME_COLOR.DARK_1,
    "lt1": MSO_THEME_COLOR.LIGHT_1,
    "dk2": MSO_THEME_COLOR.DARK_2,
    "lt2": MSO_THEME_COLOR.LIGHT_2,
    "accent1": MSO_THEME_COLOR.ACCENT_1,
    "accent2": MSO_THEME_COLOR.ACCENT_2,
    "accent3": MSO_THEME_COLOR.ACCENT_3,
    "accent4": MSO_THEME_COLOR.ACCENT_4,
    "accent5": MSO_THEME_COLOR.ACCENT_5,
    "accent6": MSO_THEME_COLOR.ACCENT_6,
    "hlink": MSO_THEME_COLOR.HYPERLINK,
    "folHlink": MSO_THEME_COLOR.FOLLOWED_HYPERLINK,
}

ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}

ANCHOR_MAP = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}

DASH_MAP = {
    "solid": MSO_LINE_DASH_STYLE.SOLID,
    "dash": MSO_LINE_DASH_STYLE.DASH,
    "dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
    "dashdot": MSO_LINE_DASH_STYLE.DASH_DOT,
}

_TOKEN_RE = re.compile(r"\[\[image:([0-9a-fA-F-]{36})")


# ---------------------------------------------------------------------------
# Colour helpers
# ---------------------------------------------------------------------------
def _apply_color(color_format, theme: dict, value) -> None:
    """Apply a resolved colour onto a python-pptx ColorFormat (font/fill/line)."""
    resolved = theme_mod.resolve_color(theme, value)
    if resolved is None:
        return
    kind, val = resolved
    if kind == "theme":
        color_format.theme_color = THEME_COLOR_MAP[val]
    else:
        color_format.rgb = RGBColor.from_string(val)


# DrawingML colour-modifier tags that python-pptx's preset ``gradient()`` bakes
# into each stop — stripped so a stop is its pure theme/hex colour and the .pptx
# gradient interpolates the same two endpoints the Pillow preview does.
_GRAD_MOD_TAGS = frozenset(
    qn(t) for t in ("a:tint", "a:shade", "a:satMod", "a:lumMod", "a:lumOff", "a:hueMod")
)


def _apply_gradient(fill, theme: dict, grad: dict, opacity: float | None = None) -> None:
    """Apply a two-stop linear gradient onto a python-pptx FillFormat.

    ``angle`` is degrees, 0 = left→right / 90 = top→bottom (matching the Pillow
    preview). python-pptx measures ``gradient_angle`` counter-clockwise, so we
    negate to land on the clockwise OOXML ``ang`` the preview uses. ``opacity``
    (<1) fades both stops so a translucent gradient panel matches the preview."""
    fill.gradient()
    stops = fill.gradient_stops
    _apply_color(stops[0].color, theme, grad.get("from", "dk2"))
    stops[0].position = 0.0
    _apply_color(stops[1].color, theme, grad.get("to", "dk1"))
    stops[1].position = 1.0
    grad_el = fill._xPr.find(qn("a:gradFill"))
    if grad_el is not None:
        # Collect colour elements before mutating — removing children mid-iter()
        # corrupts the live lxml iterator and skips the second stop.
        colours = [c for c in grad_el.iter()
                   if c.tag in (qn("a:schemeClr"), qn("a:srgbClr"))]
        for clr in colours:
            for child in list(clr):
                if child.tag in _GRAD_MOD_TAGS:
                    clr.remove(child)
            if opacity is not None and opacity < 1:
                alpha = clr.makeelement(qn("a:alpha"), {"val": str(int(max(0.0, opacity) * 100000))})
                clr.append(alpha)
    try:
        fill.gradient_angle = -float(grad.get("angle", 90.0))
    except (ValueError, TypeError, NotImplementedError):
        pass


# ---------------------------------------------------------------------------
# Text
# ---------------------------------------------------------------------------
def _render_text_body(text_frame, body: dict, theme: dict, *, default_class: str | None = None) -> None:
    text_frame.word_wrap = bool(body.get("wrap", True))
    text_frame.vertical_anchor = ANCHOR_MAP.get(body.get("valign", "top"), MSO_ANCHOR.TOP)
    cls = body.get("class") or default_class
    base = theme_mod.base_text_style(theme, cls)

    paragraphs = body.get("paragraphs") or []
    text_frame.clear()  # leaves a single empty paragraph
    for pi, para in enumerate(paragraphs):
        p = text_frame.paragraphs[0] if pi == 0 else text_frame.add_paragraph()
        p.alignment = ALIGN_MAP.get(para.get("align", "left"), PP_ALIGN.LEFT)
        if para.get("space_after") is not None:
            p.space_after = Pt(para["space_after"])
        level = int(para.get("level", 0) or 0)
        if level:
            p.level = level
        for run in para.get("runs") or []:
            r = p.add_run()
            r.text = run.get("t", "")
            style = theme_mod.merge_run(base, run)
            r.font.name = _pptx_font(theme, style.get("font"))
            if style.get("size"):
                r.font.size = Pt(style["size"])
            r.font.bold = bool(style.get("bold"))
            r.font.italic = bool(style.get("italic"))
            if style.get("underline"):
                r.font.underline = True
            _apply_color(r.font.color, theme, style.get("color"))
        # Bullet handling is a poke (last, so it stays schema-valid).
        if para.get("bullet"):
            pokes.set_bullet_char(p, theme.get("bullet", {}).get("char", "•"))
        else:
            pokes.set_no_bullet(p)


# ---------------------------------------------------------------------------
# Elements
# ---------------------------------------------------------------------------
def _pt_box(el: dict):
    return Pt(el["x"]), Pt(el["y"]), Pt(el["w"]), Pt(el["h"])


def _add_text(slide, el, theme):
    x, y, w, h = _pt_box(el)
    tb = slide.shapes.add_textbox(x, y, w, h)
    body = {
        "class": el.get("class"),
        "paragraphs": el.get("paragraphs"),
        "valign": el.get("valign", "top"),
        "wrap": el.get("wrap", True),
    }
    _render_text_body(tb.text_frame, body, theme)
    if el.get("rotation"):
        tb.rotation = el["rotation"]


def _add_harvey(slide, el, theme):
    """A Harvey ball, rasterised via the Pillow engine and embedded as a picture
    so it pixel-matches the preview. (The OOXML PIE autoshape fills its wedge on
    the opposite side from PIL's pieslice, so a native pie mis-renders in
    PowerPoint — a PNG sidesteps the whole angle-convention mismatch.)"""
    from chat.slides.pillow_render import render_harvey_png

    rgbs = _chart_ramp_rgb(theme, [el.get("fill") or "dk2"])
    rgb = rgbs[0] if rgbs else RGBColor(0x1F, 0x3D, 0x30)
    hexs = str(rgb)
    color = tuple(int(hexs[i:i + 2], 16) for i in (0, 2, 4))
    size_px = max(24, int(min(el.get("w", 18), el.get("h", 18)) * 3))
    png = render_harvey_png(color, el.get("value"), size_px)
    x, y, w, h = _pt_box(el)
    slide.shapes.add_picture(BytesIO(png), x, y, w, h)


def _add_shape(slide, el, theme, warnings):
    box = theme.get("boxes", {}).get(el["box"]) if el.get("box") else None
    box = box or {}
    shape_name = el.get("shape") or box.get("shape") or "rect"
    if shape_name == "harvey":
        _add_harvey(slide, el, theme)
        return
    mso = SHAPE_MAP.get(shape_name)
    if mso is None:
        warnings.append(f"unknown shape '{shape_name}' -> rectangle")
        mso = MSO_SHAPE.RECTANGLE
    x, y, w, h = _pt_box(el)
    shp = slide.shapes.add_shape(mso, x, y, w, h)

    fill_v = el.get("fill") if el.get("fill") is not None else box.get("fill")
    grad = el.get("gradient") or box.get("gradient")
    if grad:
        _apply_gradient(shp.fill, theme, grad, opacity=el.get("opacity"))
    elif fill_v is not None:
        shp.fill.solid()
        _apply_color(shp.fill.fore_color, theme, fill_v)
        opacity = el.get("opacity")
        if opacity is not None and opacity < 1:
            pokes.set_shape_fill_alpha(shp, opacity)
    else:
        shp.fill.background()

    line_spec = el.get("line")
    if line_spec:
        if line_spec.get("color"):
            _apply_color(shp.line.color, theme, line_spec["color"])
        if line_spec.get("w"):
            shp.line.width = Pt(line_spec["w"])
        if line_spec.get("dash"):
            shp.line.dash_style = DASH_MAP.get(line_spec["dash"], MSO_LINE_DASH_STYLE.SOLID)
    else:
        shp.line.fill.background()

    text = el.get("text")
    if text:
        default_class = box.get("class")
        # A box's text_color is the default run colour unless the text overrides.
        if box.get("text_color") and not _text_has_color(text):
            text = dict(text)
            text.setdefault("_default_color", box["text_color"])
        _render_shape_text(shp.text_frame, text, theme, default_class, box.get("text_color"))
    if el.get("rotation"):
        shp.rotation = el["rotation"]


def _text_has_color(text: dict) -> bool:
    for para in text.get("paragraphs") or []:
        for run in para.get("runs") or []:
            if run.get("color"):
                return True
    return False


def _render_shape_text(text_frame, text, theme, default_class, default_color):
    # Reuse the text-body renderer but inject the box's default colour when the
    # runs don't set one, so a themed box's label is legible on its fill.
    body = {
        "class": text.get("class"),
        "paragraphs": _inject_default_color(text.get("paragraphs"), default_color),
        "valign": text.get("valign", "middle"),
        "wrap": text.get("wrap", True),
    }
    _render_text_body(text_frame, body, theme, default_class=default_class)


def _inject_default_color(paragraphs, default_color):
    if not default_color or not paragraphs:
        return paragraphs
    out = []
    for para in paragraphs:
        p = dict(para)
        p["runs"] = [
            (run if run.get("color") else {**run, "color": default_color})
            for run in para.get("runs") or []
        ]
        out.append(p)
    return out


def _add_image(slide, el, theme, resolver, warnings):
    x, y, w, h = _pt_box(el)
    token = el.get("token", "")
    data = resolver(token) if (resolver and token) else None
    if not data:
        # A reserved slot (empty token, or a made-up slug like "demo-mockup" the
        # model uses to reserve space) reads as an intentional placeholder; only a
        # real-but-missing UUID asset is an actual problem worth a warning.
        from chat.slides.schema import _UUID_RE, image_placeholder_caption

        inner = token.split("image:", 1)[-1].rstrip("]").strip() if token else ""
        if inner and _UUID_RE.match(inner):
            warnings.append(f"image unavailable: {token[:60]}")
        ph = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        ph.fill.solid()
        ph.fill.fore_color.rgb = RGBColor(0xEC, 0xEF, 0xE9)  # lt2 sage
        ph.line.color.rgb = RGBColor(0xC7, 0xCF, 0xC4)
        tf = ph.text_frame
        tf.text = image_placeholder_caption(token)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(13)
        run.font.name = _pptx_font(theme, "data")
        run.font.color.rgb = RGBColor(0x8A, 0x94, 0x8B)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        return
    img_bytes, _ct = data
    stream = BytesIO(img_bytes)
    fit = el.get("fit", "contain")
    if fit == "stretch":
        pic = slide.shapes.add_picture(stream, x, y, width=w, height=h)
    else:
        pic = slide.shapes.add_picture(stream, x, y)  # native size first
        nat_w, nat_h = pic.width, pic.height
        if nat_w and nat_h:
            box_w, box_h = int(w), int(h)
            if fit == "cover":
                scale = max(box_w / nat_w, box_h / nat_h)
                disp_w, disp_h = nat_w * scale, nat_h * scale
                crop_x = (disp_w - box_w) / disp_w / 2 if disp_w > box_w else 0
                crop_y = (disp_h - box_h) / disp_h / 2 if disp_h > box_h else 0
                pic.crop_left = pic.crop_right = crop_x
                pic.crop_top = pic.crop_bottom = crop_y
                pic.left, pic.top, pic.width, pic.height = int(x), int(y), box_w, box_h
            else:  # contain
                scale = min(box_w / nat_w, box_h / nat_h)
                new_w, new_h = int(nat_w * scale), int(nat_h * scale)
                pic.width, pic.height = new_w, new_h
                pic.left = int(x) + (box_w - new_w) // 2
                pic.top = int(y) + (box_h - new_h) // 2
    opacity = el.get("opacity")
    if opacity is not None and opacity < 1:
        pokes.set_picture_alpha(pic, opacity)


def _add_table(slide, el, theme, warnings):
    rows = el.get("rows") or []
    n_rows = len(rows)
    n_cols = max((len(r) for r in rows), default=0)
    if n_rows == 0 or n_cols == 0:
        warnings.append("empty table skipped")
        return
    x, y, w, h = _pt_box(el)
    gf = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = gf.table
    # Neutralise the default style so our explicit fills render as specified.
    pokes.set_table_style(table, pokes.NO_STYLE_NO_GRID)
    table.first_row = False
    table.horz_banding = False

    col_widths = el.get("col_widths")
    if col_widths and len(col_widths) == n_cols:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Pt(cw)

    tbl_theme = theme.get("table", {})
    header = bool(el.get("header", True))
    banding = bool(el.get("banding", True))
    text_class = tbl_theme.get("text_class", "data")

    for ri, row in enumerate(rows):
        is_header = header and ri == 0
        is_band = banding and not is_header and (ri % 2 == 0)
        for ci in range(n_cols):
            spec = row[ci] if ci < len(row) else {}
            _render_cell(table.cell(ri, ci), spec or {}, theme, tbl_theme, text_class, is_header, is_band)

    # Content-sized rows: never stretch a small table to fill the requested h
    # (that balloons the header and detaches the data row). PowerPoint treats the
    # row height as a minimum and grows it for wrapped content.
    data_size = theme.get("text_styles", {}).get(text_class, {}).get("size", 12)
    row_h = int(Pt(max(26, data_size * 2.2)))
    for row_obj in table.rows:
        row_obj.height = row_h
    gf.height = row_h * n_rows


def _render_cell(cell, spec, theme, tbl_theme, text_class, is_header, is_band):
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = ALIGN_MAP.get(spec.get("align", "left"), PP_ALIGN.LEFT)
    r = p.add_run()
    r.text = spec.get("t", "")
    cls = spec.get("class") or text_class
    style = theme_mod.base_text_style(theme, cls)
    r.font.name = _pptx_font(theme, style.get("font"))
    # An explicit per-cell size wins over the class default (lets a dense table
    # shrink its text to fit); otherwise fall back to the text class's size.
    cell_size = spec.get("size") or style.get("size")
    if cell_size:
        r.font.size = Pt(cell_size)
    bold = spec.get("b")
    r.font.bold = bool(is_header if bold is None else bold)
    if spec.get("i") is not None:
        r.font.italic = bool(spec["i"])
    color = spec.get("color") or (tbl_theme.get("header_color") if is_header else style.get("color"))
    _apply_color(r.font.color, theme, color)

    fill_v = spec.get("fill")
    if fill_v is None:
        if is_header:
            fill_v = tbl_theme.get("header_fill")
        elif is_band:
            fill_v = tbl_theme.get("band_fill")
    if fill_v:
        cell.fill.solid()
        _apply_color(cell.fill.fore_color, theme, fill_v)
    else:
        cell.fill.background()


def _add_line(slide, el, theme):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Pt(el["x1"]), Pt(el["y1"]), Pt(el["x2"]), Pt(el["y2"])
    )
    lf = conn.line
    if el.get("color"):
        _apply_color(lf.color, theme, el["color"])
    lf.width = Pt(el.get("w") or 1)
    if el.get("dash"):
        lf.dash_style = DASH_MAP.get(el["dash"], MSO_LINE_DASH_STYLE.SOLID)
    pokes.set_connector_arrows(lf, el.get("arrow", "none"))


def _add_network(slide, el, theme, warnings):
    """A node-link diagram: edges as connectors, nodes as ovals, labels as
    textboxes. One JSON element expands to many native shapes (there's no
    per-slide shape cap on the .pptx side — the cap is on the authoring JSON)."""
    ox, oy = el.get("x", 0), el.get("y", 0)
    nodes = el.get("nodes") or []
    n = len(nodes)

    e_color, e_w = el.get("edge_color", "accent2"), el.get("edge_w", 1.0)
    for edge in el.get("edges") or []:
        a, b = edge.get("a"), edge.get("b")
        if not (isinstance(a, int) and isinstance(b, int) and 0 <= a < n and 0 <= b < n):
            continue
        x1, y1 = ox + nodes[a]["x"], oy + nodes[a]["y"]
        x2, y2 = ox + nodes[b]["x"], oy + nodes[b]["y"]
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        _apply_color(conn.line.color, theme, edge.get("color") or e_color)
        conn.line.width = Pt(edge.get("w") or e_w)
        if edge.get("dash"):
            conn.line.dash_style = DASH_MAP.get(edge["dash"], MSO_LINE_DASH_STYLE.SOLID)

    n_color = el.get("node_color", "accent2")
    n_r = el.get("node_r", 5.0)
    lbl_color = el.get("label_color", "dk1")
    lbl_size = el.get("label_size", 12.0)
    for nd in nodes:
        cx, cy = ox + nd["x"], oy + nd["y"]
        emph = bool(nd.get("emphasis"))
        base_r = nd["r"] if nd.get("r") is not None else n_r
        r = base_r * (1.6 if emph else 1.0)
        if r > 0:
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(cx - r), Pt(cy - r), Pt(2 * r), Pt(2 * r))
            dot.fill.solid()
            _apply_color(dot.fill.fore_color, theme, nd.get("color") or n_color)
            dot.line.fill.background()
            dot.shadow.inherit = False
        label = nd.get("label") or ""
        lpos = nd.get("label_pos", "r")
        if not label or lpos == "none":
            continue
        fs = nd["size"] if nd.get("size") is not None else lbl_size * (1.25 if emph else 1.0)
        lh, lw = fs * 1.5, 180.0
        gap = (r if r > 0 else 0) + 4.0
        if lpos == "l":
            lx, ly, al = cx - gap - lw, cy - lh / 2, "right"
        elif lpos == "t":
            lx, ly, al = cx - lw / 2, cy - gap - lh, "center"
        elif lpos == "b":
            lx, ly, al = cx - lw / 2, cy + gap, "center"
        elif lpos == "c":
            lx, ly, al = cx - lw / 2, cy - lh / 2, "center"
        else:  # "r"
            lx, ly, al = cx + gap, cy - lh / 2, "left"
        tb = slide.shapes.add_textbox(Pt(lx), Pt(ly), Pt(lw), Pt(lh))
        body = {"valign": "middle", "wrap": False, "paragraphs": [
            {"align": al, "runs": [{"t": label, "size": fs, "b": emph, "color": lbl_color}]}]}
        _render_text_body(tb.text_frame, body, theme)


def _chart_ramp_rgb(theme, names):
    """Resolve a list of colour refs to RGBColor for chart series/points."""
    out = []
    for name in names or []:
        resolved = theme_mod.resolve_color(theme, name)
        if resolved is None:
            continue
        kind, val = resolved
        hexv = theme["colors"].get(val) if kind == "theme" else val
        if isinstance(hexv, str):
            out.append(RGBColor.from_string(hexv.lstrip("#")))
    return out


def _add_waterfall(slide, el, theme, warnings):
    """A native, editable waterfall/bridge via the stacked-column spacer trick:
    an invisible ``base`` series lifts each floating bar, and separate
    Decrease/Increase/Total series carry the colour. python-pptx has no native
    waterfall type, so this is the standard faithful approximation."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    from chat.slides.pillow_render import _waterfall_bars

    series = el.get("series") or []
    values = list(series[0].get("values") or []) if series else []
    if not values:
        warnings.append("empty chart skipped")
        return
    n = len(values)
    cats = list(el.get("categories") or [])
    cats = cats[:n] + [str(i + 1) for i in range(len(cats), n)]  # pad/truncate to n
    bars, _edges = _waterfall_bars(values, el.get("totals") or [])

    base, dec, inc, tot = [], [], [], []
    for lo_i, hi_i, role, _delta in bars:
        height = hi_i - lo_i
        base.append(lo_i if role != "total" else 0.0)
        dec.append(height if role == "down" else 0.0)
        inc.append(height if role == "up" else 0.0)
        tot.append(height if role == "total" else 0.0)

    data = CategoryChartData()
    data.categories = cats
    data.add_series("", tuple(base))          # spacer — made transparent below
    data.add_series("Decrease", tuple(dec))
    data.add_series("Increase", tuple(inc))
    data.add_series("Total", tuple(tot))

    x, y, w, h = _pt_box(el)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, x, y, w, h, data).chart

    if el.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = el["title"]
    else:
        chart.has_title = False
    chart.has_legend = False  # the spacer series would clutter a legend

    def _one(name):
        got = _chart_ramp_rgb(theme, [name])
        return got[0] if got else None

    palette = {1: _one("danger"), 2: _one("success"), 3: _one("dk2")}
    try:
        sers = list(chart.series)
        sers[0].format.fill.background()  # transparent spacer
        for idx, color in palette.items():
            if color is not None and idx < len(sers):
                sers[idx].format.fill.solid()
                sers[idx].format.fill.fore_color.rgb = color
    except Exception:  # noqa: BLE001 — colour styling is best-effort
        logger.debug("waterfall colour styling failed", exc_info=True)


def _add_icon(slide, el, theme, warnings):
    """Rasterise a procedural icon to a theme-tinted transparent PNG and embed it
    (icons have no native OOXML form; a crisp PNG travels reliably in any .pptx)."""
    from chat.slides import icons

    rgbs = _chart_ramp_rgb(theme, [el.get("color") or "dk2"])
    rgb = rgbs[0] if rgbs else RGBColor(0x1F, 0x3D, 0x30)
    hexs = str(rgb)
    color = tuple(int(hexs[i:i + 2], 16) for i in (0, 2, 4))
    size_px = max(48, int(min(el.get("w", 24), el.get("h", 24)) * 2))
    png = icons.render_icon_png(el.get("name") or "check", size_px, color)
    if png is None:
        warnings.append(f"unknown icon '{el.get('name')}' skipped")
        return
    x, y, w, h = _pt_box(el)
    pic = slide.shapes.add_picture(BytesIO(png), x, y, w, h)
    opacity = el.get("opacity")
    if opacity is not None and opacity < 1:
        pokes.set_picture_alpha(pic, opacity)
    if el.get("rotation"):
        pic.rotation = el["rotation"]


def _add_combo(slide, el, theme, warnings):
    """A bar+line combo (dual-axis). python-pptx can't build a native combo, so
    render it via the Pillow chart engine and embed it as a crisp picture — it
    pixel-matches the on-screen preview."""
    from chat.slides.pillow_render import render_chart_png

    if not (el.get("series") or []):
        warnings.append("empty chart skipped")
        return
    png = render_chart_png(theme, el)
    x, y, w, h = _pt_box(el)
    slide.shapes.add_picture(BytesIO(png), x, y, w, h)


def _add_funnel(slide, el, theme, warnings):
    """A conversion funnel: centered rectangles of decreasing width, one per stage,
    each labelled with its stage + value. python-pptx has no funnel type."""
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    from chat.slides.pillow_render import _fmt_num

    series = el.get("series") or []
    vals = [max(0.0, float(v)) for v in (series[0].get("values") or [])
            if isinstance(v, (int, float))] if series else []
    if not vals:
        warnings.append("empty chart skipped")
        return
    cats = el.get("categories") or []
    n = len(vals)
    vmax = max(vals) or 1.0
    ramp = _chart_ramp_rgb(theme, el.get("colors") or theme.get("colors", {}).get("chart_ramp") or [])
    if not ramp:
        ramp = [RGBColor(0x2E, 0x6B, 0x52)]
    white = RGBColor(0xFF, 0xFF, 0xFF)

    x, y, w, h = el["x"], el["y"], el["w"], el["h"]  # points
    cx = x + w / 2
    gap = 4.0
    band = (h - gap * (n - 1)) / n
    for i, v in enumerate(vals):
        wd = max(10.0, (v / vmax) * w * 0.9)
        y0 = y + i * (band + gap)
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(cx - wd / 2), Pt(y0), Pt(wd), Pt(band))
        rect.fill.solid()
        rect.fill.fore_color.rgb = ramp[i % len(ramp)]
        rect.line.fill.background()
        name = cats[i] if i < len(cats) else ""
        tf = rect.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = (f"{name}  {_fmt_num(v)}").strip() if name else _fmt_num(v)
        r.font.size = Pt(9)
        r.font.color.rgb = white


def _add_marimekko(slide, el, theme, warnings):
    """A Marimekko / mosaic: variable-width 100%-stacked columns, drawn as
    rectangles (python-pptx has no native type). Column width = the size
    dimension (`widths`, else the column total); height = share of the stack."""
    from pptx.enum.text import PP_ALIGN

    from chat.slides.pillow_render import _marimekko_columns

    series = el.get("series") or []
    n_vals = max((len(s.get("values") or []) for s in series), default=0)
    cats = list(el.get("categories") or [str(i + 1) for i in range(n_vals)])
    if not series or not cats:
        warnings.append("empty chart skipped")
        return
    n_cat = len(cats)
    ws, totals = _marimekko_columns(series, cats, el.get("widths"))
    w_sum = sum(ws) or 1.0

    ramp = _chart_ramp_rgb(theme, el.get("colors") or theme.get("colors", {}).get("chart_ramp") or [])
    if not ramp:
        ramp = [RGBColor(0x2E, 0x6B, 0x52)]

    def _one(name, default):
        got = _chart_ramp_rgb(theme, [name])
        return got[0] if got else default

    txt_rgb = _one("dk1", RGBColor(0x20, 0x20, 0x20))
    grid_rgb = _one("accent3", RGBColor(0xC0, 0xC0, 0xC0))
    white = RGBColor(0xFF, 0xFF, 0xFF)

    x, y, w, h = _pt_box(el)
    multi = len(series) > 1
    title_h = Pt(20) if el.get("title") else 0
    gutter_l, gutter_b = Pt(26), Pt(16)
    legend_h = Pt(16) if multi else 0
    ax, ay = x + gutter_l, y + title_h
    aw, ah = w - gutter_l, h - title_h - gutter_b - legend_h

    def _label(bx, by, bw, bh, text, size, color, align=PP_ALIGN.CENTER):
        tb = slide.shapes.add_textbox(bx, by, bw, bh)
        tb.text_frame.word_wrap = False
        p = tb.text_frame.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = color
        return tb

    if el.get("title"):
        t = _label(x, y, w, title_h, el["title"], 12, txt_rgb)
        t.text_frame.paragraphs[0].runs[0].font.bold = True

    for pct in (0, 50, 100):
        gy = ay + ah - int(ah * pct / 100)
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ax, gy, ax + aw, gy)
        conn.line.color.rgb = grid_rgb
        conn.line.width = Pt(0.5)
        _label(x, gy - Pt(6), gutter_l - Pt(3), Pt(12), f"{pct}%", 7, txt_rgb, PP_ALIGN.RIGHT)

    gap = Pt(2)
    usable = aw - gap * (n_cat - 1)
    cx = ax
    for ci in range(n_cat):
        col_w = int(usable * (ws[ci] / w_sum))
        total = totals[ci] or 1.0
        y_bot = ay + ah
        for si, s in enumerate(series):
            vals = s.get("values") or []
            v = vals[ci] if ci < len(vals) else 0
            if not isinstance(v, (int, float)) or v <= 0:
                continue
            seg_h = int((v / total) * ah)
            y_top = y_bot - seg_h
            rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, y_top, max(col_w, 1), max(seg_h, 1))
            rect.fill.solid()
            rect.fill.fore_color.rgb = ramp[si % len(ramp)]
            rect.line.fill.background()
            if el.get("value_labels") and seg_h > Pt(12) and col_w > Pt(24):
                p = rect.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = f"{v / total * 100:.0f}%"
                r.font.size = Pt(8)
                r.font.color.rgb = white
            y_bot = y_top
        _label(cx, ay + ah, max(col_w, Pt(10)), gutter_b, cats[ci], 8, txt_rgb)
        cx += col_w + gap

    if multi:
        ly = y + h - legend_h
        lx = ax
        for si, s in enumerate(series):
            sw = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, lx, ly + Pt(3), Pt(8), Pt(8))
            sw.fill.solid()
            sw.fill.fore_color.rgb = ramp[si % len(ramp)]
            sw.line.fill.background()
            _label(lx + Pt(11), ly, Pt(90), legend_h, s.get("name", ""), 8, txt_rgb, PP_ALIGN.LEFT)
            lx += Pt(104)


def _add_chart(slide, el, theme, warnings):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    if el.get("chart") == "waterfall":
        _add_waterfall(slide, el, theme, warnings)
        return
    if el.get("chart") == "marimekko":
        _add_marimekko(slide, el, theme, warnings)
        return
    if el.get("chart") == "funnel":
        _add_funnel(slide, el, theme, warnings)
        return
    if el.get("chart") == "combo":
        _add_combo(slide, el, theme, warnings)
        return

    kind_map = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
    }
    # Stacked variants (a composition of a total). Only column/bar/area stack.
    stacked_map = {
        "column": XL_CHART_TYPE.COLUMN_STACKED,
        "bar": XL_CHART_TYPE.BAR_STACKED,
        "area": XL_CHART_TYPE.AREA_STACKED,
    }
    kind = el.get("chart", "column")
    if el.get("stacked") and kind in stacked_map:
        xl = stacked_map[kind]
    else:
        xl = kind_map.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
    series = el.get("series") or []
    if not series:
        warnings.append("empty chart skipped")
        return
    n_vals = max((len(s.get("values") or []) for s in series), default=0)
    cats = el.get("categories") or [str(i + 1) for i in range(n_vals)]

    data = CategoryChartData()
    data.categories = cats
    for s in series:
        vals = list(s.get("values") or [])
        vals += [None] * (len(cats) - len(vals))  # pad short series
        data.add_series(s.get("name", ""), tuple(vals[: len(cats)]))

    x, y, w, h = _pt_box(el)
    chart = slide.shapes.add_chart(xl, x, y, w, h, data).chart

    if el.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = el["title"]
    else:
        chart.has_title = False

    show_legend = (kind in ("pie", "doughnut") or len(series) > 1) and el.get("legend", True)
    chart.has_legend = show_legend
    if show_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    ramp = _chart_ramp_rgb(theme, el.get("colors") or theme.get("colors", {}).get("chart_ramp") or [])
    if ramp:
        try:
            if kind in ("pie", "doughnut"):
                pts = chart.plots[0].series[0].points
                for i, pt in enumerate(pts):
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = ramp[i % len(ramp)]
            else:
                for i, ser in enumerate(chart.series):
                    color = ramp[i % len(ramp)]
                    if kind in ("line",):
                        ser.format.line.color.rgb = color
                    else:
                        ser.format.fill.solid()
                        ser.format.fill.fore_color.rgb = color
        except Exception:  # noqa: BLE001 — colour styling is best-effort
            logger.debug("chart colour styling failed", exc_info=True)

    # Per-bar highlighting: colour each point of a single-series column/bar.
    pt_colors = _chart_ramp_rgb(theme, el.get("point_colors") or [])
    if pt_colors and kind in ("column", "bar") and len(series) == 1:
        try:
            pts = chart.plots[0].series[0].points
            for i, pt in enumerate(pts):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = pt_colors[i % len(pt_colors)]
        except Exception:  # noqa: BLE001 — colour styling is best-effort
            logger.debug("per-point chart colour failed", exc_info=True)

    if el.get("value_labels"):
        try:
            chart.plots[0].has_data_labels = True
        except Exception:  # noqa: BLE001
            pass

    # KPI-ring: a headline figure centred in the doughnut hole.
    if kind == "doughnut" and el.get("center_label"):
        try:
            _add_doughnut_center(slide, el, theme)
        except Exception:  # noqa: BLE001
            logger.debug("doughnut center label failed", exc_info=True)


def _add_doughnut_center(slide, el, theme):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    x, y, w, h = _pt_box(el)
    side = min(w, h) * 0.62
    tb = slide.shapes.add_textbox(x + (w - side) / 2, y + (h - side) / 2, side, side)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    label = el["center_label"]
    run.text = label
    run.font.bold = True
    # Match the preview's prominence, but shrink so a long label still fits the hole.
    mn = min(el.get("w", 100), el.get("h", 100))
    hole = el.get("hole") or 0.55
    hole_w = mn * 0.9 * hole
    size_pt = min(mn * hole * 0.30, hole_w * 1.5 / max(1, len(label)))
    run.font.size = Pt(max(12, int(size_pt)))
    _apply_color(run.font.color, theme, "dk2")


def _render_element(slide, el, theme, resolver, warnings):
    etype = el.get("type")
    try:
        if etype == "text":
            _add_text(slide, el, theme)
        elif etype == "shape":
            _add_shape(slide, el, theme, warnings)
        elif etype == "image":
            _add_image(slide, el, theme, resolver, warnings)
        elif etype == "table":
            _add_table(slide, el, theme, warnings)
        elif etype == "line":
            _add_line(slide, el, theme)
        elif etype == "chart":
            _add_chart(slide, el, theme, warnings)
        elif etype == "icon":
            _add_icon(slide, el, theme, warnings)
        elif etype == "network":
            _add_network(slide, el, theme, warnings)
        else:
            warnings.append(f"unknown element type '{etype}' skipped")
    except Exception as exc:  # noqa: BLE001 — one bad element must not fail the deck
        logger.warning("slide element render failed (%s): %s", etype, exc, exc_info=True)
        warnings.append(f"element '{etype}' failed to render")


# ---------------------------------------------------------------------------
# Slide chrome
# ---------------------------------------------------------------------------
def _apply_slide_bg(slide, theme, value):
    fill = slide.background.fill
    fill.solid()
    _apply_color(fill.fore_color, theme, value)


def _apply_bg_gradient(slide, theme, grad, prs):
    """Full-bleed gradient rectangle, drawn before the slide's elements."""
    W, H = prs.slide_width, prs.slide_height
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _apply_gradient(shp.fill, theme, grad)
    shp.line.fill.background()
    shp.shadow.inherit = False


def _apply_bg_image(slide, sdict, theme, resolver, prs, warnings):
    """Full-bleed background image (cover-cropped) + optional legibility scrim,
    drawn before the slide's elements so they sit on top."""
    W, H = prs.slide_width, prs.slide_height
    tok = sdict.get("bg_image")
    if tok and resolver:
        data = resolver(tok)
        if data:
            try:
                pic = slide.shapes.add_picture(BytesIO(data[0]), 0, 0)
                nw, nh = pic.width, pic.height
                if nw and nh:
                    scale = max(W / nw, H / nh)
                    dw, dh = nw * scale, nh * scale
                    cx = (dw - W) / dw / 2 if dw > W else 0
                    cy = (dh - H) / dh / 2 if dh > H else 0
                    pic.crop_left = pic.crop_right = cx
                    pic.crop_top = pic.crop_bottom = cy
                    pic.left, pic.top, pic.width, pic.height = 0, 0, int(W), int(H)
            except Exception:  # noqa: BLE001
                warnings.append("background image failed to render")
    scrim = sdict.get("bg_scrim")
    if scrim:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
        shp.fill.solid()
        _apply_color(shp.fill.fore_color, theme, scrim.get("color", "dk1"))
        shp.line.fill.background()
        pokes.set_shape_fill_alpha(shp, float(scrim.get("opacity", 0.4)))


def _stamp_footer(slide, theme, page_num, total):
    footer = theme.get("footer", {})
    # Optional footer logo.
    logo_name = footer.get("logo_asset")
    if logo_name:
        logo_path = ASSETS_DIR / logo_name
        if logo_path.exists():
            try:
                slide.shapes.add_picture(
                    str(logo_path), Pt(footer["x"]), Pt(footer["y"]),
                    height=Pt(footer.get("h", 18)),
                )
            except Exception:  # noqa: BLE001
                logger.warning("footer logo failed: %s", logo_name)
    # Optional footer text.
    if footer.get("text"):
        tb = slide.shapes.add_textbox(
            Pt(footer["x"] + (footer.get("h", 18) if logo_name else 0)),
            Pt(footer["y"]), Pt(footer.get("w", 200)), Pt(footer.get("h", 18)),
        )
        p = tb.text_frame.paragraphs[0]
        p.alignment = ALIGN_MAP.get(footer.get("align", "left"), PP_ALIGN.LEFT)
        r = p.add_run()
        r.text = footer["text"]
        r.font.size = Pt(footer.get("size", 9))
        r.font.name = _pptx_font(theme, "data")
        _apply_color(r.font.color, theme, footer.get("color", "dk2"))
    # Page number.
    pn = theme.get("page_number", {})
    tb = slide.shapes.add_textbox(Pt(pn.get("x", 900)), Pt(pn.get("y", 512)), Pt(pn.get("w", 36)), Pt(pn.get("h", 18)))
    p = tb.text_frame.paragraphs[0]
    p.alignment = ALIGN_MAP.get(pn.get("align", "right"), PP_ALIGN.RIGHT)
    r = p.add_run()
    r.text = str(page_num)
    r.font.size = Pt(pn.get("size", 9))
    r.font.name = _pptx_font(theme, pn.get("font", "data"))
    _apply_color(r.font.color, theme, pn.get("color", "dk2"))


# ---------------------------------------------------------------------------
# Presentation assembly
# ---------------------------------------------------------------------------
def _open_template(base_template: str) -> Presentation:
    path = TEMPLATE_DIR / f"{base_template}.pptx"
    if path.exists():
        try:
            return Presentation(str(path))
        except Exception:  # noqa: BLE001
            logger.warning("template %s unreadable; using python-pptx default", base_template)
    return Presentation()


def _blank_layout(prs: Presentation):
    # The stock template's layout[6] is "Blank"; guard for custom templates.
    layouts = prs.slide_layouts
    return layouts[6] if len(layouts) > 6 else layouts[-1]


def build_deck_pptx(
    deck: dict,
    *,
    base_template: str = "wilfred_default",
    only_slide_ids=None,
    image_resolver=None,
) -> tuple[bytes, list[str]]:
    """Render a deck dict to ``.pptx`` bytes. Returns ``(bytes, warnings)``.

    ``only_slide_ids`` renders just those slides (for the partial render cache);
    page numbers still reflect each slide's position in the FULL deck so partial
    renders match. ``image_resolver(token) -> (bytes, content_type) | None``
    supplies image bytes; a missing image degrades to a placeholder box.
    """
    warnings: list[str] = []
    theme = theme_mod.resolve_theme(deck)
    prs = _open_template(base_template)
    size = deck.get("size") or {}
    prs.slide_width = Pt(size.get("w", 960))
    prs.slide_height = Pt(size.get("h", 540))
    blank = _blank_layout(prs)

    slides = deck.get("slides") or []
    total = len(slides)
    only = set(only_slide_ids) if only_slide_ids else None

    for idx, sdict in enumerate(slides):
        if only is not None and sdict.get("id") not in only:
            continue
        slide = prs.slides.add_slide(blank)
        if sdict.get("bg"):
            _apply_slide_bg(slide, theme, sdict["bg"])
        if sdict.get("bg_gradient"):
            _apply_bg_gradient(slide, theme, sdict["bg_gradient"], prs)
        if sdict.get("bg_image") or sdict.get("bg_scrim"):
            _apply_bg_image(slide, sdict, theme, image_resolver, prs, warnings)
        for el in sdict.get("elements") or []:
            _render_element(slide, el, theme, image_resolver, warnings)
        if not sdict.get("skip_footer"):
            _stamp_footer(slide, theme, idx + 1, total)
        if sdict.get("notes"):
            slide.notes_slide.notes_text_frame.text = sdict["notes"]

    out = BytesIO()
    prs.save(out)
    return out.getvalue(), warnings


# ---------------------------------------------------------------------------
# Django-aware wrapper
# ---------------------------------------------------------------------------
def build_pptx(slide_set, *, only_slide_ids=None) -> tuple[bytes, list[str]]:
    """Build ``.pptx`` bytes for a ``SlideSet`` model instance."""
    deck = slide_set.content or {}
    resolver = make_image_resolver(slide_set)
    return build_deck_pptx(
        deck,
        base_template=slide_set.base_template or "wilfred_default",
        only_slide_ids=only_slide_ids,
        image_resolver=resolver,
    )


def make_image_resolver(slide_set):
    """A resolver bound to a deck's owner: only assets owned by this deck's
    thread or the deck itself resolve, so a model-supplied token can't leak
    another user's asset bytes. Caches per (uuid) within one build.
    """
    from django.db.models import Q

    from chat.assets import image_asset_source
    from chat.models import Asset

    cache: dict[str, tuple[bytes, str] | None] = {}

    def resolve(token: str):
        match = _TOKEN_RE.search(token or "")
        if not match:
            return None
        aid = match.group(1).lower()
        if aid in cache:
            return cache[aid]
        result = None
        try:
            asset = Asset.objects.filter(
                Q(pk=aid)
                & (Q(slide_set=slide_set) | Q(thread_id=slide_set.thread_id))
            ).first()
            if asset is not None:
                source, ct = image_asset_source(asset)
                if source is not None:
                    with source.open("rb") as fh:
                        result = (fh.read(), ct or "image/png")
        except Exception:  # noqa: BLE001
            logger.warning("image token resolve failed: %s", aid, exc_info=True)
        cache[aid] = result
        return result

    return resolve
