"""Tests for ``.pptx`` font embedding (chat.slides.pptx_fonts)."""
from __future__ import annotations

import io
import zipfile

from django.test import TestCase


def _small_deck() -> dict:
    return {"version": 1, "size": {"w": 960, "h": 540}, "slides": [{"id": "s1", "elements": []}]}


def _caladea_faces():
    from core.fonts import bundled_resolution

    return bundled_resolution("Caladea").faces


class PptxFontEmbedding(TestCase):
    def test_embed_injects_parts_rels_and_lst_and_reopens(self):
        from chat.slides.pptx_build import build_deck_pptx
        from chat.slides.pptx_fonts import embed_fonts

        pptx_bytes, _ = build_deck_pptx(_small_deck())
        out = embed_fonts(pptx_bytes, [{"typeface": "Acme Brand", "faces": _caladea_faces()}])

        z = zipfile.ZipFile(io.BytesIO(out))
        names = z.namelist()
        self.assertTrue(any(n.startswith("ppt/fonts/font") and n.endswith(".fntdata") for n in names))
        ct = z.read("[Content_Types].xml").decode("utf-8")
        self.assertIn("fntdata", ct)
        self.assertIn("x-fontdata", ct)
        pres = z.read("ppt/presentation.xml").decode("utf-8")
        self.assertIn("embedTrueTypeFonts", pres)
        self.assertIn("embeddedFontLst", pres)
        self.assertIn("Acme Brand", pres)
        rels = z.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
        self.assertIn("relationships/font", rels)

        # The modified package must still be a valid .pptx python-pptx can re-open.
        from pptx import Presentation

        Presentation(io.BytesIO(out))

    def test_build_embeds_nonbundled_only(self):
        from chat.slides.pptx_build import build_deck_pptx

        faces = _caladea_faces()
        out_c, _ = build_deck_pptx(_small_deck(), font_bundle={
            "Acme Brand": {"name": "Acme Brand", "bundled": False, "faces": faces}})
        self.assertTrue(any(n.startswith("ppt/fonts/")
                            for n in zipfile.ZipFile(io.BytesIO(out_c)).namelist()))

        out_b, _ = build_deck_pptx(_small_deck(), font_bundle={
            "Caladea": {"name": "Caladea", "bundled": True, "faces": faces}})
        self.assertFalse(any(n.startswith("ppt/fonts/")
                             for n in zipfile.ZipFile(io.BytesIO(out_b)).namelist()))

    def test_woff_faces_are_not_embedded(self):
        from chat.slides.pptx_build import _fonts_to_embed
        from core.fonts import FontFace

        woff = [FontFace(weight=400, style="normal", data=b"wOFFxxxx", fmt="woff")]
        self.assertEqual(_fonts_to_embed({"W": {"bundled": False, "faces": woff}}), [])

    def test_no_font_bundle_leaves_pptx_unembedded(self):
        from chat.slides.pptx_build import build_deck_pptx

        out, _ = build_deck_pptx(_small_deck())
        self.assertFalse(any(n.startswith("ppt/fonts/")
                             for n in zipfile.ZipFile(io.BytesIO(out)).namelist()))
