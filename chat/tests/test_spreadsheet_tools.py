"""Tests for the spreadsheet chat tools (document_read_sheet /
document_view_sheet): manifest-driven reads, gates, tile attach flow, caching
and caps. Rendering itself is patched (WeasyPrint is unavailable on Windows).
"""
import io
import json
import tempfile
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings

from documents.models import DataRoom, DataRoomDocument, DataRoomDocumentVersion
from llm.types.context import RunContext

User = get_user_model()

_MEDIA = tempfile.mkdtemp()


def _workbook_bytes():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Projects"
    ws.append(["Number", "Project name", "Status"])
    ws.append([202501007, "Ultra-Performance AM", "ACTIVE"])
    ws.append([202501008, "Beta project", "TERMINATED"])
    ws.append([202501009, "Gamma project", "ACTIVE"])
    ws.row_dimensions[4].hidden = True
    ws.auto_filter.ref = "A1:C4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _manifest_for(xlsx_bytes: bytes) -> dict:
    from pathlib import Path

    from documents.services.spreadsheets.headers import detect_headers
    from documents.services.spreadsheets.manifest import build_manifest
    from documents.services.spreadsheets.mesh import plan_sheet_mesh
    from documents.services.spreadsheets.model import load_workbook_model

    with tempfile.TemporaryDirectory() as tmpdir:
        path = Path(tmpdir) / "wb.xlsx"
        path.write_bytes(xlsx_bytes)
        model = load_workbook_model(path)
    headers = {s.index: detect_headers(s) for s in model.sheets}
    plans = {s.index: plan_sheet_mesh(s, headers[s.index].header_row) for s in model.sheets}
    return build_manifest(model, headers, headers, plans, {})


@override_settings(MEDIA_ROOT=_MEDIA)
class SpreadsheetToolTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="sheet@test.com", password="pw")
        self.room = DataRoom.objects.create(name="R", slug="r-sheet", created_by=self.user)
        self.xlsx = _workbook_bytes()
        self.doc = DataRoomDocument.objects.create(
            data_room=self.room, uploaded_by=self.user,
            original_filename="projects.xlsx", doc_index=1,
            status=DataRoomDocument.Status.READY,
        )
        self.version = DataRoomDocumentVersion.objects.create(
            document=self.doc, parser_type="openpyxl",
            native_blob=ContentFile(self.xlsx, name="projects.xlsx"),
            processing_metadata=_manifest_for(self.xlsx),
        )
        self.doc.current_version = self.version
        self.doc.save(update_fields=["current_version"])

    def _read_tool(self, data_room_ids=None):
        from chat.spreadsheet_tools import DocumentReadSheetTool

        tool = DocumentReadSheetTool()
        tool.set_context(RunContext.create(
            user_id=self.user.pk, data_room_ids=data_room_ids or [self.room.pk],
        ))
        return tool

    def _view_tool(self, data_room_ids=None):
        from chat.spreadsheet_tools import DocumentViewSheetTool

        tool = DocumentViewSheetTool()
        tool.set_context(RunContext.create(
            user_id=self.user.pk, data_room_ids=data_room_ids or [self.room.pk],
        ))
        return tool


class ReadSheetToolTests(SpreadsheetToolTestCase):
    def test_overview(self):
        result = json.loads(self._read_tool()._run(1))
        self.assertEqual(result["file"], "projects.xlsx")
        sheet = result["sheets"][0]
        self.assertEqual(sheet["name"], "Projects")
        self.assertEqual(sheet["header_row"], 1)
        self.assertEqual(sheet["columns"]["A"], "Number")
        self.assertTrue(sheet["tiles"])
        tool = self._read_tool()
        self.assertEqual(tool.end_label_for_result(result), "Listed 1 sheet")

    def test_find_returns_exact_cells_with_labels_and_tiles(self):
        result = json.loads(self._read_tool()._run(1, find="ultra-performance"))
        self.assertEqual(result["match_count"], 1)
        match = result["matches"][0]
        self.assertEqual(match["cell"], "B2")
        self.assertEqual(match["label"], "Project name")
        self.assertEqual(match["value"], "Ultra-Performance AM")
        self.assertEqual(match["tile"], "x0y0")
        tool = self._read_tool()
        self.assertEqual(tool.end_label_for_result(result), "Found 1 match in the spreadsheet")

    def test_find_flags_hidden_rows(self):
        result = json.loads(self._read_tool()._run(1, find="Gamma"))
        self.assertEqual(result["match_count"], 1)
        self.assertTrue(result["matches"][0].get("hidden"))

    @override_settings(XLSX_FIND_MAX_RESULTS=1)
    def test_find_truncates_at_cap(self):
        result = json.loads(self._read_tool()._run(1, find="ACTIVE"))
        self.assertEqual(result["match_count"], 1)
        self.assertTrue(result["truncated"])

    def test_find_unknown_sheet_errors(self):
        result = json.loads(self._read_tool()._run(1, find="x", sheet="Nope"))
        self.assertIn("No sheet", result["error"])

    def test_rows_read_with_labels(self):
        result = json.loads(self._read_tool()._run(1, rows="2-3"))
        self.assertEqual(result["sheet"], "Projects")
        self.assertEqual(result["columns"]["B"], "Project name")
        self.assertEqual(len(result["rows"]), 2)
        self.assertEqual(result["rows"][0]["row"], 2)
        self.assertEqual(result["rows"][0]["cells"]["A"], "202501007")
        self.assertIn("tiles", result)

    def test_rows_marks_header_and_hidden(self):
        result = json.loads(self._read_tool()._run(1, rows="1-4"))
        by_row = {r["row"]: r for r in result["rows"]}
        self.assertTrue(by_row[1].get("is_header"))
        self.assertTrue(by_row[4].get("hidden"))

    def test_cols_filter(self):
        result = json.loads(self._read_tool()._run(1, rows="2-3", cols="A,C"))
        self.assertEqual(set(result["rows"][0]["cells"]), {"A", "C"})

    @override_settings(XLSX_READ_MAX_ROWS=2)
    def test_rows_cap_truncates(self):
        result = json.loads(self._read_tool()._run(1, cols="A"))
        self.assertEqual(len(result["rows"]), 2)
        self.assertTrue(result["truncated"])

    def test_cell_read(self):
        result = json.loads(self._read_tool()._run(1, cell="C3"))
        self.assertEqual(result["value"], "TERMINATED")
        self.assertEqual(result["label"], "Status")
        self.assertIn("Project name (B)", result["row_values"])

    def test_access_denied_for_foreign_room(self):
        other = User.objects.create_user(email="other@test.com", password="pw")
        room2 = DataRoom.objects.create(name="R2", slug="r2-sheet", created_by=other)
        result = json.loads(self._read_tool([room2.pk])._run(1))
        self.assertIn("error", result)

    def test_scanning_document_reads_as_not_found(self):
        self.doc.status = DataRoomDocument.Status.SCANNING
        self.doc.save(update_fields=["status"])
        result = json.loads(self._read_tool()._run(1))
        self.assertIn("No document with index 1", result["error"])

    def test_partial_quarantine_refuses(self):
        # Stricter than document_read on purpose: direct workbook reads bypass
        # per-chunk quarantine.
        self.version.is_partially_quarantined = True
        self.version.save(update_fields=["is_partially_quarantined"])
        result = json.loads(self._read_tool()._run(1))
        self.assertIn("quarantined", result["error"])

    def test_full_quarantine_refuses(self):
        self.version.is_quarantined = True
        self.version.save(update_fields=["is_quarantined"])
        result = json.loads(self._read_tool()._run(1))
        self.assertIn("quarantined", result["error"])

    def test_non_spreadsheet_document_errors(self):
        self.version.processing_metadata = None
        self.version.save(update_fields=["processing_metadata"])
        result = json.loads(self._read_tool()._run(1))
        self.assertIn("not a spreadsheet", result["error"])


def _fake_render_band(sheet, plan, band_x, y_limit=None):
    from documents.services.spreadsheets.render import band_row_groups

    n = len(band_row_groups(sheet, plan, band_x))
    if y_limit is not None:
        n = min(n, y_limit)
    return [(b"\x89PNG-fake-tile", plan.tile_w, plan.tile_h)] * n


class ViewSheetToolTests(SpreadsheetToolTestCase):
    def test_default_view_attaches_top_tiles_with_tokens(self):
        tool = self._view_tool()
        with patch("documents.services.spreadsheets.tiles.render_band", side_effect=_fake_render_band):
            result = tool._run(1)
        self.assertIn("Attached tile x0y0", result)
        self.assertIn("columns A–C", result)
        pending = tool.context.pending_image_assets
        self.assertEqual(len(pending), 1)
        self.assertEqual(pending[0]["media_type"], "image/png")
        self.assertTrue(pending[0]["asset_id"].startswith("[[image:"))
        self.assertIn("tile x0y0", pending[0]["description"])

    def test_rows_selects_covering_tiles(self):
        tool = self._view_tool()
        with patch("documents.services.spreadsheets.tiles.render_band", side_effect=_fake_render_band):
            result = tool._run(1, rows="2-3")
        self.assertIn("Attached tile x0y0", result)
        self.assertEqual(len(tool.context.pending_image_assets), 1)

    def test_second_view_hits_tile_cache(self):
        tool = self._view_tool()
        with patch(
            "documents.services.spreadsheets.tiles.render_band", side_effect=_fake_render_band
        ) as render:
            tool._run(1)
            self.assertEqual(render.call_count, 1)
            tool2 = self._view_tool()
            tool2._run(1)
            self.assertEqual(render.call_count, 1)  # cached Asset reused
        self.assertEqual(len(tool2.context.pending_image_assets), 1)

    @override_settings(XLSX_MAX_TILES_PER_VIEW=1)
    def test_tile_cap_reports_excess(self):
        tool = self._view_tool()
        with patch("documents.services.spreadsheets.tiles.render_band", side_effect=_fake_render_band):
            result = tool._run(1, tiles=["x0y0", "x0y0"])
        self.assertIn("showing the first 1", result)
        self.assertEqual(len(tool.context.pending_image_assets), 1)

    def test_invalid_tile_and_missing_tile_reported(self):
        tool = self._view_tool()
        with patch("documents.services.spreadsheets.tiles.render_band", side_effect=_fake_render_band):
            result = tool._run(1, tiles=["bogus", "x9y9"])
        self.assertIn("Invalid tile reference", result)
        self.assertIn("does not exist", result)
        self.assertEqual(tool.context.pending_image_assets, [])

    def test_render_unavailable_degrades_to_message(self):
        tool = self._view_tool()
        with patch(
            "documents.services.spreadsheets.tiles.render_band",
            side_effect=RuntimeError("Spreadsheet tile rendering is unavailable in this environment."),
        ):
            result = tool._run(1)
        self.assertIn("unavailable", result)
        self.assertEqual(tool.context.pending_image_assets, [])

    def test_quarantine_refuses(self):
        self.version.is_partially_quarantined = True
        self.version.save(update_fields=["is_partially_quarantined"])
        result = self._view_tool()._run(1)
        self.assertIn("quarantined", result)


class CollectDocImagesTileExclusionTests(SpreadsheetToolTestCase):
    def test_tiles_are_excluded_from_document_view_image(self):
        from chat.models import Asset
        from chat.tools import _collect_doc_images
        from documents.services.spreadsheets.tiles import TILE_ALT_MARKER, store_tile

        store_tile(self.version, self.doc, 0, 0, 0, b"\x89PNG-tile", 1568, 768)
        embedded = Asset(
            version=self.version, content_type="image/png", size_bytes=4,
            description="an embedded chart", created_by=self.user,
        )
        embedded.blob.save(f"{embedded.id}.png", ContentFile(b"\x89PNG-embedded"), save=True)

        images = _collect_doc_images(self.doc)
        descriptions = [d for _, _, d in images]
        self.assertEqual(descriptions, ["an embedded chart"])
        self.assertTrue(
            Asset.objects.filter(version=self.version, alt_text=TILE_ALT_MARKER).exists()
        )
